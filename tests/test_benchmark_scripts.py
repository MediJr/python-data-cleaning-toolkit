"""The benchmark generator and the independent verifier keep working (small dataset).

The verifier must accept correct outputs and reject altered ones. measure.py and
run_benchmark.py (Windows timing) are not exercised here; the CLI is called directly.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from data_cleaning.cli import main

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts" / "benchmark"
ROWS = 2_000


@pytest.fixture(scope="module")
def workdir(tmp_path_factory) -> Path:
    work = tmp_path_factory.mktemp("benchmark")
    subprocess.run([sys.executable, str(SCRIPTS / "generate_data.py"), str(work), "--rows", str(ROWS)], check=True, capture_output=True)
    manifest = json.loads((work / "data" / "manifest.json").read_text(encoding="utf-8"))
    results = []
    commands = {
        "customers": [manifest["customers"]["file"], "customers_clean.csv", "--config", str(ROOT / "config/customers_config.json")],
        "orders": [manifest["orders"]["file"], "orders_clean.xlsx", "--sheet", "Orders", "--config", str(ROOT / "config/orders_config.json")],
    }
    for name, (source, output, *options) in commands.items():
        out_dir = work / "runs" / name / "run1"
        code = main(["clean", str(work / "data" / source), str(out_dir / output), *options])
        results.append({"workflow": name, "run": 1, "exit_code": code})
    (work / "results.json").write_text(json.dumps(results), encoding="utf-8")
    return work


def _verify(work: Path, workflow: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, "-I", str(SCRIPTS / "verify_outputs.py"), str(work), workflow, "1"], capture_output=True, text=True
    )


def test_generator_is_deterministic(workdir, tmp_path):
    subprocess.run([sys.executable, str(SCRIPTS / "generate_data.py"), str(tmp_path), "--rows", str(ROWS)], check=True, capture_output=True)
    first = json.loads((workdir / "data/manifest.json").read_text(encoding="utf-8"))
    second = json.loads((tmp_path / "data/manifest.json").read_text(encoding="utf-8"))
    # openpyxl stores the save time inside .xlsx files, so only the workbook's bytes (and its
    # fingerprint) may differ; everything else, including the CSV bytes, must be identical.
    for manifest in (first, second):
        manifest["orders"].pop("sha256")
    assert first == second
    assert (workdir / "data" / f"customers_{ROWS}.csv").read_bytes() == (tmp_path / "data" / f"customers_{ROWS}.csv").read_bytes()
    workbooks = [load_workbook(d / "data" / f"orders_{ROWS}.xlsx", read_only=True)["Orders"] for d in (workdir, tmp_path)]
    assert [list(r) for r in workbooks[0].iter_rows(values_only=True)] == [list(r) for r in workbooks[1].iter_rows(values_only=True)]
    assert first["customers"]["rows"] == first["orders"]["rows"] == ROWS


@pytest.mark.parametrize("workflow", ["customers", "orders"])
def test_verifier_accepts_the_real_outputs(workdir, workflow):
    completed = _verify(workdir, workflow)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert ", 0 failed" in completed.stdout


def test_verifier_rejects_an_altered_output(workdir):
    cleaned = workdir / "runs/customers/run1/customers_clean.csv"
    original = cleaned.read_text(encoding="utf-8")
    lines = original.split("\n")
    lines[10] = lines[10].lstrip("0")  # drop the leading zeros of one customer ID
    cleaned.write_text("\n".join(lines), encoding="utf-8")
    try:
        completed = _verify(workdir, "customers")
    finally:
        cleaned.write_text(original, encoding="utf-8")
    assert completed.returncode == 1
    assert "row-by-row mismatches {'customer_id': 1}" in completed.stdout
