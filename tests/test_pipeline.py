"""End-to-end tests of the cleaning workflow and of report accuracy."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from data_cleaning import pipeline
from data_cleaning.config import validate_config
from data_cleaning.errors import ExportError, InputFileError, OutputExistsError
from data_cleaning.io import read_table
from data_cleaning.pipeline import EXIT_ERROR, EXIT_OK, EXIT_QUALITY_ISSUES, plan_outputs, run_cleaning

RAW_CSV = (
    "Customer ID, City ,Signup Date,Amount,Status\n"
    '007,  new york ,15/03/2024,"1,200.50",active\n'  # line 2
    "008,NEW YORK,03/04/2024,N/A,Active\n"  # line 3: ambiguous date, missing token
    "009,Paris,2024-02-30,abc,closed\n"  # line 4: invalid date and number
    "009,Paris,2024-02-30,abc,closed\n"  # line 5: exact duplicate of line 4
    "010,Lyon,2024-05-01,-5,unknown\n"  # line 6: invalid status (quarantine), negative amount
)

CONFIG = {
    "column_names": {"style": "snake_case"},
    "missing_values": {"tokens": ["N/A"]},
    "text": [{"columns": ["city"], "collapse_whitespace": True}],
    "categories": [
        {"column": "city", "mapping": {"new york": "New York", "NEW YORK": "New York"}},
        {"column": "status", "mapping": {"Active": "active"}},
    ],
    "dates": [{"column": "signup_date", "formats": ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]}],
    "numbers": [{"column": "amount", "thousands": ","}],
    "duplicates": {"remove_exact": True, "key_columns": ["customer_id"]},
    "validation": [
        {"rule": "allowed_values", "column": "status", "values": ["active", "closed"], "action": "quarantine", "id": "status"},
        {"rule": "range", "column": "amount", "min": 0, "id": "amount"},
    ],
}


@pytest.fixture
def raw_csv(tmp_path: Path) -> Path:
    path = tmp_path / "raw.csv"
    path.write_text(RAW_CSV, encoding="utf-8", newline="")
    return path


def _run(raw: Path, out: Path, config: dict | None = None, **kwargs):
    return run_cleaning(raw, out, config=validate_config(config if config is not None else CONFIG), **kwargs)


def test_full_run_produces_expected_files_values_and_exit_code(raw_csv, tmp_path):
    before = raw_csv.read_bytes()
    result = _run(raw_csv, tmp_path / "out" / "clean.csv")
    out = tmp_path / "out"
    assert result.exit_code == EXIT_QUALITY_ISSUES
    assert raw_csv.read_bytes() == before
    assert sorted(p.name for p in out.iterdir()) == [
        "clean.csv", "clean_quarantine.csv", "clean_removed_rows.csv", "clean_report.json", "clean_report.md",
    ]
    assert (out / "clean.csv").read_text(encoding="utf-8") == (
        "customer_id,city,signup_date,amount,status\n"
        "007,New York,2024-03-15,1200.5,active\n"
        "008,New York,03/04/2024,,active\n"
        "009,Paris,2024-02-30,abc,closed\n"
    )
    assert (out / "clean_quarantine.csv").read_text(encoding="utf-8") == (
        "customer_id,city,signup_date,amount,status,quarantine_reason\n"
        "010,Lyon,2024-05-01,-5,unknown,[status] 'unknown' is not an allowed value\n"
    )
    assert (out / "clean_removed_rows.csv").read_text(encoding="utf-8") == (
        "customer_id,city,signup_date,amount,status,removal_reason\n"
        "009,Paris,2024-02-30,abc,closed,exact duplicate of row 4\n"
    )


def test_report_numbers_match_the_files_actually_written(raw_csv, tmp_path):
    result = _run(raw_csv, tmp_path / "clean.csv")
    report = json.loads((tmp_path / "clean_report.json").read_text(encoding="utf-8"))
    assert report == json.loads(json.dumps(result.report, default=str))

    assert report["execution"] == {"status": "success", "error": None, "error_type": None}
    assert report["data_quality"]["status"] == "issues_found"
    dims = report["dimensions"]
    assert dims["initial"] == {"rows": 5, "columns": 5}
    assert dims["final"] == {"rows": 3, "columns": 5}
    assert dims["rows_removed"] == 1 and dims["rows_quarantined"] == 1
    assert dims["initial"]["rows"] == dims["final"]["rows"] + dims["rows_removed"] + dims["rows_quarantined"]
    for name, file_name in (("cleaned", "clean.csv"), ("quarantine", "clean_quarantine.csv"), ("removed_rows", "clean_removed_rows.csv")):
        reopened = read_table(tmp_path / file_name).data
        assert report["outputs"][name]["rows"] == len(reopened)
        assert report["outputs"][name]["columns"] == reopened.shape[1]
        assert report["outputs"][name]["verified"] is True

    assert report["input"]["unchanged"] is True and report["input"]["sha256_before"] == report["input"]["sha256_after"]
    assert report["columns"]["renamed"] == {
        "Customer ID": "customer_id", " City ": "city", "Signup Date": "signup_date", "Amount": "amount", "Status": "status",
    }
    assert report["columns"]["cells_differing_from_input"] == {"city": 2, "signup_date": 1, "amount": 2, "status": 1}
    assert {m["column"]: (m["before"], m["after"]) for m in report["missing_values"]}["amount"] == (1, 1)

    failures = [(f["source_row"], f["column"], f["in_cleaned_output"]) for f in report["conversion_failures"]]
    assert failures == [
        (3, "signup_date", True), (4, "signup_date", True), (5, "signup_date", False), (4, "amount", True), (5, "amount", False),
    ]
    assert [(a["type"], a["count"]) for a in report["anomalies"]] == [
        ("conversion_failure", 3), ("validation_failed", 1), ("validation_failed", 2),
    ]
    rules = {r["id"]: (r["status"], r["failures"]) for r in report["validation"]["rules"]}
    assert rules == {"status": ("failed", 1), "amount": ("failed", 2)}
    assert report["duplicates"]["exact"]["duplicate_rows"] == 1 and report["duplicates"]["exact"]["removed"] == 1


def test_markdown_report_states_both_statuses_and_never_claims_clean_data(raw_csv, tmp_path):
    _run(raw_csv, tmp_path / "clean.csv")
    markdown = (tmp_path / "clean_report.md").read_text(encoding="utf-8")
    assert "| Execution status | SUCCESS |" in markdown
    assert "| Data-quality status | ISSUES FOUND |" in markdown
    assert "not a certification" in markdown
    assert "| 3 | signup_date | 03/04/2024 | ambiguous date" in markdown
    assert "| ' City ' | 'city' |" in markdown
    assert "data is clean" not in markdown.lower()


def test_rules_that_pass_give_accepted_status_and_exit_0(tmp_path):
    raw = tmp_path / "ok.csv"
    raw.write_text("id,qty\n001,5\n002,7\n", encoding="utf-8")
    config = {"numbers": [{"column": "qty"}], "validation": [{"rule": "range", "column": "qty", "min": 0}]}
    result = _run(raw, tmp_path / "out.csv", config)
    assert result.exit_code == EXIT_OK
    assert result.report["data_quality"]["status"] == "accepted"
    assert "configured" in result.report["data_quality"]["summary"]


def test_without_rules_quality_is_not_claimed(tmp_path):
    raw = tmp_path / "ok.csv"
    raw.write_text("id\n001\n", encoding="utf-8")
    result = run_cleaning(raw, tmp_path / "out.csv")
    assert result.exit_code == EXIT_OK
    assert result.report["data_quality"]["status"] == "no_rules_configured"
    assert (tmp_path / "out.csv").read_text(encoding="utf-8") == "id\n001\n"


def test_defaults_detect_but_do_not_remove_duplicates(tmp_path):
    raw = tmp_path / "dups.csv"
    raw.write_text("a,b\n1,x\n1,x\n", encoding="utf-8")
    result = run_cleaning(raw, tmp_path / "out.csv")
    assert result.exit_code == EXIT_QUALITY_ISSUES
    assert len(read_table(tmp_path / "out.csv").data) == 2
    assert result.report["anomalies"][0]["type"] == "exact_duplicates_kept"
    assert not (tmp_path / "out_removed_rows.csv").exists()
    flagged = run_cleaning(raw, tmp_path / "out2.csv", remove_exact_duplicates=True)
    assert len(read_table(tmp_path / "out2.csv").data) == 1
    assert flagged.report["overrides"] == {"duplicates.remove_exact": True}


def test_key_duplicate_removal_keeps_last_and_saves_removed_rows(tmp_path):
    raw = tmp_path / "orders.csv"
    raw.write_text("order,amount\nA,1\nA,2\nB,3\n", encoding="utf-8")
    config = {"duplicates": {"key_columns": ["order"], "remove_key_duplicates": True, "keep": "last"}}
    result = _run(raw, tmp_path / "out.csv", config)
    assert (tmp_path / "out.csv").read_text(encoding="utf-8") == "order,amount\nA,2\nB,3\n"
    assert (tmp_path / "out_removed_rows.csv").read_text(encoding="utf-8") == (
        "order,amount,removal_reason\nA,1,duplicate of row 3 on key ['order'] (keep=last)\n"
    )
    assert any("differing values" in w for w in result.report["warnings"])


def test_xlsx_to_xlsx_with_sheet_selection_keeps_codes_as_text_and_numbers_numeric(tmp_path):
    source = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active.title = "Notes"
    workbook.active.append(["ignore"])
    sheet = workbook.create_sheet("Orders")
    sheet.append(["Code", "Qty"])
    sheet.append(["007", "1,000"])
    sheet.append(["010", 25])
    workbook.save(source)
    result = _run(source, tmp_path / "clean.xlsx", {"numbers": [{"column": "Qty", "thousands": ","}]}, sheet="Orders")
    assert result.exit_code == EXIT_OK
    cleaned = load_workbook(tmp_path / "clean.xlsx")["cleaned"]
    assert [[c.value for c in row] for row in cleaned.iter_rows()] == [["Code", "Qty"], ["007", 1000], ["010", 25]]
    assert cleaned["A2"].data_type == "s" and cleaned["B2"].data_type == "n"
    assert result.report["input"]["sheet"] == "Orders"


def test_csv_to_xlsx_and_back_preserves_values(raw_csv, tmp_path):
    _run(raw_csv, tmp_path / "a.xlsx")
    _run(tmp_path / "a.xlsx", tmp_path / "b.csv", {})
    assert read_table(tmp_path / "b.csv").data.values.tolist() == read_table(tmp_path / "a.xlsx").data.values.tolist()


def test_repeated_runs_are_deterministic(raw_csv, tmp_path):
    first = _run(raw_csv, tmp_path / "one" / "clean.csv")
    second = _run(raw_csv, tmp_path / "two" / "clean.csv")
    for name in ("clean.csv", "clean_quarantine.csv", "clean_removed_rows.csv"):
        assert (tmp_path / "one" / name).read_bytes() == (tmp_path / "two" / name).read_bytes()

    def comparable(report):
        report = json.loads(json.dumps(report, default=str))
        for key in ("generated_at", "outputs", "planned_outputs"):
            report.pop(key)
        return report

    assert comparable(first.report) == comparable(second.report)


# --------------------------------------------------------------------------- safety and failures


def test_existing_outputs_stop_the_run_before_anything_is_written(raw_csv, tmp_path):
    (tmp_path / "clean_report.md").write_text("keep me", encoding="utf-8")
    (tmp_path / "clean.csv").write_text("keep me too", encoding="utf-8")
    with pytest.raises(ExportError, match="(?s)cannot be written.*clean.csv.*clean_report.md"):
        _run(raw_csv, tmp_path / "clean.csv")
    assert (tmp_path / "clean.csv").read_text(encoding="utf-8") == "keep me too"
    assert not (tmp_path / "clean_quarantine.csv").exists()
    assert (tmp_path / "clean_report.md").read_text(encoding="utf-8") == "keep me"


def test_overwrite_replaces_previous_outputs(raw_csv, tmp_path):
    _run(raw_csv, tmp_path / "clean.csv")
    result = _run(raw_csv, tmp_path / "clean.csv", overwrite=True)
    assert result.exit_code == EXIT_QUALITY_ISSUES and result.report["overrides"] == {"overwrite": True}


def test_input_can_never_be_overwritten(raw_csv):
    before = raw_csv.read_bytes()
    with pytest.raises(ExportError, match="input file"):
        _run(raw_csv, raw_csv, overwrite=True)
    assert raw_csv.read_bytes() == before


def test_report_cannot_target_the_input_file(tmp_path):
    raw = tmp_path / "data_report.json"
    raw.write_text("x", encoding="utf-8")
    assert plan_outputs(tmp_path / "data.csv").report_json == raw
    with pytest.raises(ExportError, match="input file"):
        run_cleaning(raw, tmp_path / "data.csv", overwrite=True)
    assert raw.read_text(encoding="utf-8") == "x"


def test_unsupported_output_type_is_rejected_before_reading(raw_csv, tmp_path):
    with pytest.raises(ExportError, match="Unsupported output type"):
        _run(raw_csv, tmp_path / "clean.json")


def test_missing_input_is_reported_without_writing_anything(tmp_path):
    with pytest.raises(InputFileError, match="not found"):
        run_cleaning(tmp_path / "missing.csv", tmp_path / "out.csv")
    assert list(tmp_path.iterdir()) == []


def test_configured_column_that_does_not_exist_fails_with_report_and_no_output(raw_csv, tmp_path):
    config = {"column_names": {"style": "snake_case"}, "numbers": [{"column": "Amount"}]}
    result = _run(raw_csv, tmp_path / "clean.csv", config)
    assert result.exit_code == EXIT_ERROR
    assert not (tmp_path / "clean.csv").exists()
    error = result.report["execution"]["error"]
    assert "numbers[0]: 'Amount' (renamed to 'amount'" in error
    assert result.report["data_quality"]["status"] == "not_assessed"
    markdown = (tmp_path / "clean_report.md").read_text(encoding="utf-8")
    assert "| Execution status | FAILED |" in markdown and "NOT WRITTEN" in markdown


def test_column_name_collision_fails_safely(tmp_path):
    raw = tmp_path / "raw.csv"
    raw.write_text("Order ID,order_id\n1,2\n", encoding="utf-8")
    result = _run(raw, tmp_path / "clean.csv", {"column_names": {"style": "snake_case"}})
    assert result.exit_code == EXIT_ERROR
    assert result.report["execution"]["error_type"] == "ColumnCollisionError"
    assert not (tmp_path / "clean.csv").exists()


def test_failure_after_main_export_is_reported_as_incomplete(raw_csv, tmp_path, monkeypatch):
    real_write = pipeline.write_table

    def failing_write(data, path, **kwargs):
        if "quarantine" in Path(path).name:
            raise ExportError("disk full (simulated)")
        return real_write(data, path, **kwargs)

    monkeypatch.setattr(pipeline, "write_table", failing_write)
    result = _run(raw_csv, tmp_path / "clean.csv")
    assert result.exit_code == EXIT_ERROR
    report = result.report
    assert report["execution"]["status"] == "failed" and "disk full" in report["execution"]["error"]
    assert list(report["outputs"]) == ["cleaned"]  # what was actually written is listed
    assert report["data_quality"]["status"] == "not_assessed"


def test_modification_of_the_input_during_a_run_is_detected(raw_csv, tmp_path, monkeypatch):
    real_process = pipeline._process

    def tampering_process(ctx, raw):
        real_process(ctx, raw)
        with raw_csv.open("a", encoding="utf-8") as handle:
            handle.write("999,X,2024-01-01,1,active\n")

    monkeypatch.setattr(pipeline, "_process", tampering_process)
    result = _run(raw_csv, tmp_path / "clean.csv")
    assert result.exit_code == EXIT_ERROR
    assert "input file changed" in result.report["execution"]["error"]
    assert result.report["input"]["unchanged"] is False


def test_unexpected_internal_error_is_reraised_after_writing_a_failure_report(raw_csv, tmp_path, monkeypatch):
    def broken_profile(*args, **kwargs):
        raise KeyError("simulated bug")

    monkeypatch.setattr(pipeline, "profile_dataframe", broken_profile)
    with pytest.raises(KeyError, match="simulated bug"):
        _run(raw_csv, tmp_path / "clean.csv")
    report = json.loads((tmp_path / "clean_report.json").read_text(encoding="utf-8"))
    assert report["execution"]["status"] == "failed"
    assert "Unexpected internal error" in report["execution"]["error"]
    assert not (tmp_path / "clean.csv").exists()


def test_stale_companion_file_is_reported_not_deleted(tmp_path):
    raw = tmp_path / "raw.csv"
    raw.write_text("a\n1\n", encoding="utf-8")
    stale = tmp_path / "out_quarantine.csv"
    stale.write_text("old", encoding="utf-8")
    result = run_cleaning(raw, tmp_path / "out.csv", overwrite=True)
    assert stale.read_text(encoding="utf-8") == "old"
    assert any("earlier run" in w for w in result.report["warnings"])


def test_existing_quarantine_file_without_overwrite_is_refused(raw_csv, tmp_path):
    (tmp_path / "clean_quarantine.csv").write_text("old", encoding="utf-8")
    with pytest.raises(OutputExistsError, match="clean_quarantine.csv"):
        _run(raw_csv, tmp_path / "clean.csv")
    assert not (tmp_path / "clean.csv").exists()


def test_uncalculated_formulas_make_the_quality_status_issues_found(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "f.xlsx"
    workbook = Workbook()
    workbook.active.append(["id", "total"])
    workbook.active.append(["001", "=1+1"])
    workbook.save(source)
    result = run_cleaning(source, tmp_path / "out.csv")
    assert result.exit_code == EXIT_QUALITY_ISSUES
    assert result.report["data_quality"]["status"] == "issues_found"
    assert result.report["anomalies"] == [
        {"type": "uncalculated_formulas", "count": 1, "description": "1 formula cell(s) in the input have no saved value and were read as empty."}
    ]
    assert result.report["input"]["uncalculated_formula_cells"] == 1
