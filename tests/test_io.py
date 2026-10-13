"""Tests for reading, writing and verifying CSV/XLSX files."""

from __future__ import annotations

import datetime as dt
import os
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from data_cleaning.errors import (
    EmptyInputError,
    ExportError,
    InputFileError,
    MalformedInputError,
    OutputExistsError,
)
from data_cleaning.io import read_table, render_text, write_table
from conftest import make_frame


# --------------------------------------------------------------------------- CSV reading


def test_csv_values_are_read_as_text_and_identifiers_keep_leading_zeros(write_text):
    path = write_text("in.csv", "id,zip,amount\n007,01234,1.50\n010,99999,0\n")
    table = read_table(path)
    assert table.data["id"].tolist() == ["007", "010"]
    assert table.data["zip"].tolist() == ["01234", "99999"]
    assert table.data["amount"].tolist() == ["1.50", "0"]
    assert all(str(dtype) == "string" for dtype in table.data.dtypes)
    assert table.data.index.tolist() == [2, 3]  # source line numbers, header on line 1


def test_csv_empty_cells_are_kept_as_empty_strings_not_converted(write_text):
    path = write_text("in.csv", "a,b\n,NULL\nN/A,0\n")
    data = read_table(path).data
    assert data["a"].tolist() == ["", "N/A"]
    assert data["b"].tolist() == ["NULL", "0"]


def test_csv_semicolon_delimiter_is_detected_and_recorded(write_text):
    path = write_text("in.csv", "name;price\nTea;1,50\nCoffee;2,00\n")
    table = read_table(path)
    assert table.delimiter == ";" and table.delimiter_detected
    assert table.data["price"].tolist() == ["1,50", "2,00"]


def test_csv_explicit_delimiter_overrides_detection(write_text):
    path = write_text("in.csv", "a|b\n1|2\n")
    table = read_table(path, delimiter="|")
    assert list(table.data.columns) == ["a", "b"] and not table.delimiter_detected


def test_csv_utf8_bom_is_removed_from_first_header(write_text):
    path = write_text("in.csv", "city,n\nZürich,1\n", encoding="utf-8-sig")
    table = read_table(path)
    assert list(table.data.columns) == ["city", "n"]
    assert table.data["city"].tolist() == ["Zürich"]
    assert table.encoding == "utf-8" and not table.warnings


def test_csv_cp1252_fallback_is_reported_as_warning(write_text):
    path = write_text("in.csv", "city,n\nMontréal,1\n", encoding="cp1252")
    table = read_table(path)
    assert table.data["city"].tolist() == ["Montréal"]
    assert table.encoding == "cp1252"
    assert any("Windows-1252" in w for w in table.warnings)


def test_csv_wrong_explicit_encoding_is_an_error(write_text):
    path = write_text("in.csv", "city\nMontréal\n", encoding="cp1252")
    with pytest.raises(MalformedInputError, match="not valid utf-8"):
        read_table(path, encoding="utf-8")


def test_csv_quoted_fields_with_delimiters_and_newlines(write_text):
    path = write_text("in.csv", 'id,note\n1,"a, b"\n2,"line1\nline2"\n3,x\n')
    data = read_table(path).data
    assert data["note"].tolist() == ["a, b", "line1\nline2", "x"]
    assert data.index.tolist() == [2, 3, 5]  # row 3 spans lines 3-4


def test_csv_row_with_too_few_fields_is_rejected_not_padded(write_text):
    path = write_text("in.csv", "a,b,c\n1,2,3\n4,5\n")
    with pytest.raises(MalformedInputError, match=r"line 3 \(2 fields\)"):
        read_table(path)


def test_csv_ambiguous_delimiter_requires_explicit_option(write_text):
    path = write_text("in.csv", "a;b,c\n1;2\n3,4,5,6\n")
    with pytest.raises(MalformedInputError, match="delimiter"):
        read_table(path)


def test_csv_row_with_too_many_fields_is_rejected(write_text):
    path = write_text("in.csv", "a,b\n1,2\n3,4,5\n", )
    with pytest.raises(MalformedInputError, match=r"line 3 \(3 fields\)"):
        read_table(path, delimiter=",")


def test_csv_broken_quoting_is_rejected(write_text):
    path = write_text("in.csv", 'a,b\n"unterminated,2\n')
    with pytest.raises(MalformedInputError):
        read_table(path, delimiter=",")


def test_csv_duplicate_header_names_are_rejected(write_text):
    path = write_text("in.csv", "id,name,name\n1,a,b\n")
    with pytest.raises(MalformedInputError, match="duplicate column names: 'name'"):
        read_table(path)


def test_csv_blank_header_gets_positional_name_with_warning(write_text):
    path = write_text("in.csv", "id,,value\n1,x,2\n")
    table = read_table(path, delimiter=",")
    assert list(table.data.columns) == ["id", "column_2", "value"]
    assert any("column_2" in w for w in table.warnings)


def test_csv_completely_empty_rows_are_skipped_and_counted(write_text):
    path = write_text("in.csv", "a,b\n1,2\n\n,\n3,4\n")
    table = read_table(path, delimiter=",")
    assert table.data["a"].tolist() == ["1", "3"]
    assert table.skipped_blank_rows == 1  # the ',' row; the physically empty line is not a record
    assert table.data.index.tolist() == [2, 5]


def test_empty_file_is_rejected(write_text):
    with pytest.raises(EmptyInputError, match="0 bytes"):
        read_table(write_text("empty.csv", ""))


def test_whitespace_only_file_is_rejected(write_text):
    with pytest.raises(EmptyInputError):
        read_table(write_text("blank.csv", "\n  \n"))


def test_header_only_file_is_rejected(write_text):
    with pytest.raises(EmptyInputError, match="no data rows"):
        read_table(write_text("header.csv", "a,b\n"))


def test_missing_path_and_directory_and_unsupported_type(tmp_path):
    with pytest.raises(InputFileError, match="not found"):
        read_table(tmp_path / "nope.csv")
    folder = tmp_path / "folder.csv"
    folder.mkdir()
    with pytest.raises(InputFileError, match="not a file"):
        read_table(folder)
    macro = tmp_path / "book.xlsm"
    macro.write_bytes(b"x")
    with pytest.raises(InputFileError, match="Macro-enabled"):
        read_table(macro)


def test_binary_content_with_csv_extension_is_rejected(write_text, tmp_path):
    path = tmp_path / "fake.csv"
    path.write_bytes(b"a,b\n\x00\x01,\x02\n")
    with pytest.raises(MalformedInputError, match="NUL"):
        read_table(path)


def test_sheet_option_is_rejected_for_csv(write_text):
    with pytest.raises(InputFileError, match="sheet"):
        read_table(write_text("in.csv", "a\n1\n"), sheet="Sheet1")


# --------------------------------------------------------------------------- XLSX reading


def _workbook(path: Path, sheets: dict[str, list[list]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


def test_xlsx_typed_cells_are_converted_to_text_without_loss(tmp_path):
    path = _workbook(
        tmp_path / "in.xlsx",
        {
            "Data": [
                ["id", "qty", "price", "when", "stamp", "flag"],
                ["00123", 5, 2.5, dt.datetime(2024, 3, 1), dt.datetime(2024, 3, 1, 14, 30), True],
            ]
        },
    )
    data = read_table(path).data
    assert data.iloc[0].tolist() == ["00123", "5", "2.5", "2024-03-01", "2024-03-01 14:30:00", "TRUE"]


def test_xlsx_sheet_selection_and_multi_sheet_warning(tmp_path):
    path = _workbook(tmp_path / "in.xlsx", {"First": [["a"], ["1"]], "Second": [["b"], ["2"]]})
    default = read_table(path)
    assert default.sheet == "First" and any("2 sheets" in w for w in default.warnings)
    chosen = read_table(path, sheet="Second")
    assert chosen.data["b"].tolist() == ["2"] and not chosen.warnings
    with pytest.raises(InputFileError, match="Available sheets: First, Second"):
        read_table(path, sheet="Missing")


def test_xlsx_empty_sheet_is_rejected(tmp_path):
    path = _workbook(tmp_path / "in.xlsx", {"Empty": []})
    with pytest.raises(EmptyInputError):
        read_table(path)


def test_xlsx_row_numbers_and_trailing_empty_columns(tmp_path):
    path = tmp_path / "in.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"] = "a", "b"
    sheet["A2"], sheet["B2"] = "x", "y"
    sheet["A4"] = "z"
    sheet["D1"].number_format = "0.00"  # formatted but empty column
    sheet["D4"].number_format = "0.00"
    workbook.save(path)
    table = read_table(path)
    assert list(table.data.columns) == ["a", "b"]
    assert table.data.index.tolist() == [2, 4]
    assert table.skipped_blank_rows == 1


def test_corrupt_xlsx_is_rejected_with_explanation(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(MalformedInputError, match="corrupt, password-protected"):
        read_table(path)


# --------------------------------------------------------------------------- writing


@pytest.fixture
def typed_frame() -> pd.DataFrame:
    frame = make_frame({"id": ["00123", "=1+1", "  padded  "], "note": ["é", None, "x"]})
    frame["qty"] = pd.array([1, None, 3], dtype="Int64")
    frame["price"] = pd.array([1.5, 0.1, None], dtype="Float64")
    return frame


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_export_round_trip_is_verified_and_lossless(tmp_path, typed_frame, extension):
    target = tmp_path / f"out.{extension}"
    result = write_table(typed_frame, target)
    assert result.verified and result.rows == 3 and result.columns == 4
    reloaded = read_table(target).data
    assert list(reloaded.columns) == ["id", "note", "qty", "price"]
    assert reloaded["id"].tolist() == ["00123", "=1+1", "  padded  "]
    assert reloaded["note"].tolist() == ["é", "", "x"]
    assert reloaded["qty"].tolist() == ["1", "", "3"]
    assert reloaded["price"].tolist() == ["1.5", "0.1", ""]
    assert not list(tmp_path.glob(".*tmp-*")), "temporary file left behind"


def test_xlsx_export_stores_numbers_as_numbers_and_formula_text_as_text(tmp_path, typed_frame):
    target = tmp_path / "out.xlsx"
    write_table(typed_frame, target)
    sheet = load_workbook(target)["cleaned"]
    assert sheet["C2"].value == 1 and sheet["C2"].data_type == "n"
    assert sheet["D2"].value == 1.5
    assert sheet["A3"].value == "=1+1" and sheet["A3"].data_type == "s"
    assert sheet["A2"].value == "00123" and sheet["A2"].data_type == "s"


def test_existing_output_is_not_overwritten_without_permission(tmp_path, typed_frame):
    target = tmp_path / "out.csv"
    target.write_text("original", encoding="utf-8")
    with pytest.raises(OutputExistsError):
        write_table(typed_frame, target)
    assert target.read_text(encoding="utf-8") == "original"
    write_table(typed_frame, target, overwrite=True)
    assert target.read_text(encoding="utf-8").startswith("id,note,qty,price")


def test_input_file_can_never_be_the_output_even_with_overwrite(tmp_path, typed_frame, write_text):
    source = write_text("in.csv", "a\n1\n")
    before = source.read_bytes()
    with pytest.raises(ExportError, match="input file"):
        write_table(typed_frame, source, overwrite=True, protected_paths=[source])
    # A different spelling of the same path is also caught.
    with pytest.raises(ExportError, match="input file"):
        write_table(typed_frame, tmp_path / "." / "in.csv", overwrite=True, protected_paths=[source])
    assert source.read_bytes() == before


def test_unsupported_output_type_and_directory_target(tmp_path, typed_frame):
    with pytest.raises(ExportError, match="Unsupported output type"):
        write_table(typed_frame, tmp_path / "out.json")
    (tmp_path / "dir.csv").mkdir()
    with pytest.raises(ExportError, match="directory"):
        write_table(typed_frame, tmp_path / "dir.csv", overwrite=True)


def test_output_encoding_failure_is_reported_and_no_file_is_left(tmp_path, typed_frame):
    target = tmp_path / "out.csv"
    with pytest.raises(ExportError, match="encode"):
        write_table(typed_frame, target, csv_encoding="ascii")
    assert not target.exists()
    assert not list(tmp_path.iterdir())


def test_control_characters_cannot_be_written_to_xlsx(tmp_path):
    frame = make_frame({"a": ["bad\x01value"]})
    with pytest.raises(ExportError, match="control character"):
        write_table(frame, tmp_path / "out.xlsx")
    assert not list(tmp_path.iterdir())


@pytest.mark.skipif(os.name != "nt", reason="Windows file locking")
def test_locked_output_file_reports_permission_error(tmp_path, typed_frame):
    target = tmp_path / "out.csv"
    target.write_text("x", encoding="utf-8")
    # On Windows an open handle (e.g. the file open in Excel) prevents replacing the file.
    with target.open("rb"):
        with pytest.raises(ExportError, match="Permission denied"):
            write_table(typed_frame, target, overwrite=True)
    assert target.read_text(encoding="utf-8") == "x"
    assert [p.name for p in tmp_path.iterdir()] == ["out.csv"]


def test_render_text_formats_numbers_canonically():
    frame = make_frame({"t": ["a", None]})
    frame["f"] = pd.array([2.0, 0.30000000000000004], dtype="Float64")
    rendered = render_text(frame)
    assert rendered["f"].tolist() == ["2", "0.30000000000000004"]
    assert rendered["t"].isna().tolist() == [False, True]


# --------------------------------------------------------------------------- formulas and export verification


def _formula_workbook(path: Path, cached: dict[str, str] | None = None) -> Path:
    """Workbook with two formula cells; ``cached`` injects saved values as Excel would."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["id", "qty", "price", "total"])
    sheet.append(["001", 2, 5, "=B2*C2"])
    sheet.append(["002", 3, 4, "=B3*C3"])
    workbook.save(path)
    if cached:
        import zipfile

        with zipfile.ZipFile(path) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        xml = members["xl/worksheets/sheet1.xml"].decode("utf-8")
        for formula, value in cached.items():
            empty_value = f"<f>{formula}</f><v />"
            assert empty_value in xml, "test fixture: unexpected openpyxl XML"
            xml = xml.replace(empty_value, f"<f>{formula}</f><v>{value}</v>")
        members["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in members.items():
                archive.writestr(name, content)
    return path


def test_uncalculated_formulas_are_read_as_empty_but_never_silently(tmp_path):
    """Regression: formulas without a saved value used to become empty cells with no warning."""
    table = read_table(_formula_workbook(tmp_path / "f.xlsx"))
    assert table.data["total"].tolist() == ["", ""]
    assert (table.formula_cells, table.uncalculated_formula_cells) == (2, 2)
    assert any("2 formula cell(s)" in w and "read as EMPTY" in w for w in table.warnings)
    assert table.describe()["uncalculated_formula_cells"] == 2


def test_formulas_with_saved_values_use_those_values(tmp_path):
    table = read_table(_formula_workbook(tmp_path / "f.xlsx", cached={"B2*C2": "10", "B3*C3": "12"}))
    assert table.data["total"].tolist() == ["10", "12"]
    assert (table.formula_cells, table.uncalculated_formula_cells) == (2, 0)
    assert any("contain formulas; their last values saved by Excel were used" in w for w in table.warnings)
    assert not any("read as EMPTY" in w for w in table.warnings)


def test_workbook_without_formulas_has_no_formula_warning(tmp_path):
    table = read_table(_workbook(tmp_path / "plain.xlsx", {"S": [["a"], ["1"]]}))
    assert (table.formula_cells, table.uncalculated_formula_cells) == (0, 0)
    assert not any("formula" in w for w in table.warnings)


def test_export_that_does_not_read_back_identically_is_rejected(tmp_path, typed_frame, monkeypatch):
    """Adversarial: if the written file does not match the data, the export must fail and leave nothing."""
    import data_cleaning.io as io_module

    real_read = io_module.read_table

    def corrupted_read(path, **kwargs):
        table = real_read(path, **kwargs)
        table.data.iat[0, 0] = "tampered"
        return table

    monkeypatch.setattr(io_module, "read_table", corrupted_read)
    target = tmp_path / "out.csv"
    with pytest.raises(ExportError, match="Verification failed: column 'id', data row 1: wrote '00123' but read back 'tampered'"):
        write_table(typed_frame, target)
    assert not target.exists()
    assert list(tmp_path.iterdir()) == []
