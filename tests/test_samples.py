"""Sample workflows: the real CLI on the synthetic datasets, checked against
expectations that were derived by hand (``data/sample/expected/`` and the
numbers below), plus a check that the committed ``reports/examples/`` still
match a fresh run.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from data_cleaning.cli import main
from data_cleaning.io import read_table

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "sample"
EXPECTED = SAMPLE / "expected"
EXAMPLES = ROOT / "reports" / "examples"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def workspace(tmp_path: Path, monkeypatch) -> Path:
    """A copy of the sample data and configs; commands run with the same relative paths as the docs."""
    shutil.copytree(ROOT / "data", tmp_path / "data")
    shutil.copytree(ROOT / "config", tmp_path / "config")
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _clean_customers(out_dir: str = "output") -> int:
    return main(["clean", "data/sample/customers_raw.csv", f"{out_dir}/customers_clean.csv", "--config", "config/customers_config.json"])


def _clean_orders(out_dir: str = "output") -> int:
    return main(
        ["clean", "data/sample/orders_raw.xlsx", f"{out_dir}/orders_clean.xlsx", "--sheet", "Orders", "--config", "config/orders_config.json"]
    )


def _report(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------- customers (CSV)


def test_customers_outputs_match_hand_derived_expectations(workspace):
    source = workspace / "data/sample/customers_raw.csv"
    digest = _sha256(source)
    assert _clean_customers() == 3
    assert _sha256(source) == digest
    for name in ("customers_clean.csv", "customers_clean_quarantine.csv", "customers_clean_removed_rows.csv"):
        assert (workspace / "output" / name).read_text(encoding="utf-8") == (EXPECTED / name).read_text(encoding="utf-8"), name


def test_customers_values_that_must_not_be_changed(workspace):
    _clean_customers()
    rows = {r["customer_id"]: r for r in read_table("output/customers_clean.csv").data.to_dict("records")}
    assert rows["000103"]["signup_date"] == "03/04/2024"  # ambiguous: DD/MM or MM/DD, not guessed
    assert rows["000105"]["signup_date"] == "2024-02-30"  # impossible date kept, reported
    assert rows["000105"]["postal_code"] == "2108"  # looks like a lost leading zero, but is not padded
    assert rows["000105"]["annual_spend"] == "-50"  # suspicious negative value reported, not changed
    assert rows["000104"]["annual_spend"] == "1.234,50"  # European format not guessed under '.' decimal rule
    assert rows["000111"]["annual_spend"] == "0"  # zero is a value, not missing
    assert rows["000102"]["full_name"] == "ben okafor"  # spaces fixed, case deliberately not changed
    assert rows["000113"]["city"] == "Bostn"  # possible typo not mapped automatically
    assert rows["000101"]["postal_code"] == "02108" and rows["000101"]["customer_id"] == "000101"


def test_customers_report_matches_the_data(workspace):
    _clean_customers()
    report = _report(workspace / "output/customers_clean_report.json")
    assert report["execution"]["status"] == "success"
    assert report["data_quality"]["status"] == "issues_found"
    assert report["data_quality"]["unresolved_anomaly_types"] == 11
    assert report["data_quality"]["unresolved_occurrences"] == 22
    assert report["dimensions"] == {
        "initial": {"rows": 15, "columns": 8},
        "final": {"rows": 11, "columns": 8},
        "rows_removed": 1,
        "rows_quarantined": 3,
        "blank_rows_skipped_on_import": 0,
    }
    assert report["profile_before"]["exact_duplicate_rows"] == 1
    assert report["profile_before"]["exact_duplicate_groups"] == [[9, 10]]
    assert report["duplicates"]["exact"]["duplicate_rows"] == 1 and report["duplicates"]["exact"]["removed"] == 1
    keys = report["duplicates"]["keys"]
    assert [g["source_rows"] for g in keys["groups"]] == [[7, 8]]
    assert keys["groups"][0]["differing_columns"] == ["email", "signup_date", "annual_spend"]
    assert keys["rows_with_missing_key"] == 1 and keys["removed"] == 0

    missing = {m["column"]: (m["before"], m["after"]) for m in report["missing_values"]}
    assert missing == {
        "customer_id": (1, 0), "full_name": (0, 0), "email": (1, 1), "city": (0, 0),
        "signup_date": (0, 0), "annual_spend": (3, 3), "status": (0, 0), "postal_code": (0, 0),
    }
    assert report["columns"]["cells_differing_from_input"] == {
        "full_name": 1, "email": 1, "city": 4, "signup_date": 2, "annual_spend": 5, "status": 2,
    }
    failures = [(f["source_row"], f["column"], f["value"], f["in_cleaned_output"]) for f in report["conversion_failures"]]
    assert failures == [
        (4, "signup_date", "03/04/2024", True),
        (6, "signup_date", "2024-02-30", True),
        (12, "signup_date", "yesterday", False),
        (5, "annual_spend", "1.234,50", True),
        (12, "annual_spend", "twelve hundred", False),
    ]
    assert "ambiguous" in report["conversion_failures"][0]["reason"]
    rules = {r["id"]: (r["status"], r["failures"]) for r in report["validation"]["rules"]}
    assert rules == {
        "required_columns": ("passed", 0),
        "id_present": ("failed", 1),
        "id_unique": ("failed", 2),
        "email_present": ("failed", 1),
        "email_format": ("failed", 1),
        "status_allowed": ("failed", 2),
        "signup_is_iso_date": ("failed", 3),
        "signup_in_range": ("failed", 4),
        "spend_non_negative": ("failed", 3),
        "postal_code_5_digits": ("failed", 1),
    }
    # 18 issues, but only 9 distinct rows (4, 5, 6, 7, 8, 11, 12, 13, 15): rules overlap.
    assert (report["validation"]["issue_count"], report["validation"]["rows_with_issues"]) == (18, 9)
    failing_rows = {(i["rule_id"], i["source_row"]) for i in report["validation"]["issues"]}
    assert ("signup_in_range", 13) in failing_rows and ("postal_code_5_digits", 6) in failing_rows
    assert any("'Bostn'" in w for w in report["warnings"])


# --------------------------------------------------------------------------- orders (XLSX)


def test_orders_outputs_match_hand_derived_expectations(workspace):
    source = workspace / "data/sample/orders_raw.xlsx"
    digest = _sha256(source)
    assert _clean_orders() == 3
    assert _sha256(source) == digest
    for name in ("orders_clean", "orders_clean_quarantine", "orders_clean_removed_rows"):
        produced = read_table(workspace / "output" / f"{name}.xlsx").data
        expected = read_table(EXPECTED / f"{name}.csv").data
        assert list(produced.columns) == list(expected.columns), name
        assert produced.values.tolist() == expected.values.tolist(), name


def test_orders_xlsx_cell_types(workspace):
    _clean_orders()
    sheet = load_workbook(workspace / "output/orders_clean.xlsx")["cleaned"]
    header = [c.value for c in sheet[1]]
    column = {name: header.index(name) for name in header}
    rows = list(sheet.iter_rows(min_row=2))
    prices = [r[column["unit_price"]] for r in rows]
    assert [c.value for c in prices] == [19.99, 5, 12.5, 1299, None, 7.5]
    assert all(c.data_type == "n" for c in prices if c.value is not None)  # fully converted -> numeric cells
    assert [r[column["customer_id"]].value for r in rows][2] == "101"  # stays text, never padded or converted
    assert all(r[column["customer_id"]].data_type == "s" for r in rows)
    assert all(r[column["quantity"]].data_type == "s" for r in rows)  # 'two' failed, so the column stays text
    assert rows[5][column["order_date"]].value == "2024-03-06 14:30:00"  # time of day not truncated


def test_orders_report_matches_the_data(workspace):
    _clean_orders()
    report = _report(workspace / "output/orders_clean_report.json")
    assert report["input"]["sheet"] == "Orders" and report["input"]["available_sheets"] == ["Read me", "Orders"]
    assert report["dimensions"]["initial"] == {"rows": 8, "columns": 7}
    assert report["dimensions"]["final"] == {"rows": 6, "columns": 7}
    assert (report["dimensions"]["rows_removed"], report["dimensions"]["rows_quarantined"]) == (1, 1)
    assert report["data_quality"]["unresolved_anomaly_types"] == 6
    assert report["data_quality"]["unresolved_occurrences"] == 10
    assert report["duplicates"]["keys"]["duplicate_key_groups"] == 0
    # Conversions run before duplicate removal, so the removed duplicate (row 6) also has a
    # failure; it is listed but not counted as unresolved because it is not in the output.
    assert [(f["source_row"], f["column"], f["in_cleaned_output"]) for f in report["conversion_failures"]] == [
        (7, "order_date", True), (9, "order_date", True), (5, "quantity", True), (6, "quantity", False),
    ]
    assert report["anomalies"][0] == {
        "type": "conversion_failure",
        "count": 3,
        "description": "3 date/number value(s) in the cleaned output could not be converted "
        "(1 more in rows that were removed or quarantined) (see conversion failures).",
    }
    assert {r["id"]: r["failures"] for r in report["validation"]["rules"]} == {
        "order_id_present": 1, "customer_id_6_digits": 1, "quantity_at_least_1": 2, "channel_allowed": 1, "order_date_is_iso_date": 2,
    }
    assert {m["column"]: (m["before"], m["after"]) for m in report["missing_values"] if m["before"] or m["after"]} == {
        "order_id": (1, 0), "unit_price": (1, 1),
    }
    assert (report["validation"]["issue_count"], report["validation"]["rows_with_issues"]) == (7, 5)  # rows 4, 5, 7, 8, 9
    assert report["columns"]["cells_differing_from_input"] == {"order_date": 2, "sku": 1, "unit_price": 2, "channel": 3}


# --------------------------------------------------------------------------- invalid inputs


@pytest.mark.parametrize(
    "file_name, message",
    [
        ("ragged_rows.csv", "line 3 (2 fields)"),
        ("duplicate_headers.csv", "duplicate column names: 'amount'"),
        ("empty.csv", "empty (0 bytes)"),
        ("header_only.csv", "no data rows"),
    ],
)
def test_invalid_inputs_are_rejected_without_output(workspace, capsys, file_name, message):
    assert main(["clean", f"data/sample/invalid/{file_name}", "output/out.csv"]) == 1
    assert message in capsys.readouterr().err
    assert not (workspace / "output").exists()


def test_column_collision_sample_fails_with_report_and_no_data_output(workspace, capsys):
    code = main(["clean", "data/sample/invalid/column_collision.csv", "output/out.csv", "--config", "config/invalid/collision_config.json"])
    assert code == 1
    report = _report(workspace / "output/out_report.json")
    assert report["execution"]["error_type"] == "ColumnCollisionError"
    assert "'Order ID' / 'order_id' -> 'order_id'" in report["execution"]["error"]
    assert sorted(p.name for p in (workspace / "output").iterdir()) == ["out_report.json", "out_report.md"]


def test_typo_in_config_is_rejected(workspace, capsys):
    assert main(["clean", "data/sample/customers_raw.csv", "output/out.csv", "--config", "config/invalid/typo_config.json"]) == 1
    assert "Unknown key(s) in configuration: missing_value" in capsys.readouterr().err
    assert not (workspace / "output").exists()


# --------------------------------------------------------------------------- committed examples and sample hygiene


def test_sales_demo_matches_the_scenario_design(workspace):
    """docs/demo-sales-cleaning.md: every planted anomaly is fixed, kept and reported, or set aside as designed."""
    code = main(["clean", "data/sample/sales_raw.csv", "output/sales_clean.csv", "--config", "config/sales_config.json"])
    assert code == 3  # success, but issues remain on purpose
    cleaned = read_table(workspace / "output/sales_clean.csv").data
    rows = {(r["order_id"], r["quantity"]): r for r in cleaned.to_dict("records")}
    assert len(cleaned) == 13
    # fixed: labels, spaces, unambiguous date, currency and thousands separators, missing tokens
    assert rows[("SO-1002", "1")]["region"] == "North" and rows[("SO-1002", "1")]["channel"] == "Online"
    assert rows[("SO-1002", "1")]["sales_rep"] == "Ana Silva" and rows[("SO-1002", "1")]["unit_price"] == "89.9"
    assert rows[("SO-1003", "5")]["order_date"] == "2026-09-15" and rows[("SO-1003", "5")]["channel"] == "Online"
    assert rows[("SO-1001", "2")]["unit_price"] == "1250" and rows[("SO-1001", "2")]["sku"] == "00123"
    assert rows[("SO-1006", "")]["quantity"] == "" and rows[("SO-1013", "7")]["region"] == ""
    # kept unchanged (and reported): ambiguous or impossible dates, unexpected number format, suspicious values
    assert rows[("SO-1004", "3")]["order_date"] == "05/09/2026"
    assert rows[("SO-1005", "1")]["order_date"] == "2026-09-31"
    assert rows[("SO-1008", "4")]["unit_price"] == "12,5" and ("SO-1008", "6") in rows  # key duplicate: both kept
    assert rows[("SO-1007", "-2")]["region"] == "East" and rows[("SO-1006", "")]["region"] == "Sth"
    assert rows[("SO-1012", "1")]["sku"] == "0123" and rows[("SO-1012", "1")]["order_date"] == "2026-10-02"
    quarantine = read_table(workspace / "output/sales_clean_quarantine.csv").data
    assert quarantine["quarantine_reason"].tolist() == [
        "[order_id_present] 'order_id' is missing",
        "[channel_allowed] 'Fax' is not an allowed value",
    ]
    removed = read_table(workspace / "output/sales_clean_removed_rows.csv").data
    assert removed["removal_reason"].tolist() == ["exact duplicate of row 11"]
    report = _report(workspace / "output/sales_clean_report.json")
    assert report["dimensions"]["initial"]["rows"] == 16 and report["dimensions"]["final"]["rows"] == 13
    assert [(f["source_row"], f["column"]) for f in report["conversion_failures"]] == [
        (5, "order_date"),
        (6, "order_date"),
        (9, "unit_price"),
    ]
    assert {r["id"]: r["failures"] for r in report["validation"]["rules"]} == {
        "order_id_present": 1,
        "channel_allowed": 1,
        "region_present": 1,
        "region_allowed": 1,
        "order_date_is_iso_date": 2,
        "order_date_in_september": 3,
        "quantity_present": 1,
        "quantity_at_least_1": 1,
        "unit_price_is_number": 1,
        "sku_5_digits": 1,
        "order_line_unique": 2,
    }
    assert (report["validation"]["issue_count"], report["validation"]["rows_with_issues"]) == (15, 10)
    assert report["input"]["sha256_before"] == report["input"]["sha256_after"] == _sha256(SAMPLE / "sales_raw.csv")


def _normalized_json(path: Path) -> dict:
    report = _report(path)
    report.pop("generated_at")
    return report


def _normalized_markdown(path: Path) -> str:
    return re.sub(r"^Generated \S+ ", "Generated <time> ", path.read_text(encoding="utf-8"), flags=re.M)


def _load_script(name: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_committed_examples_match_a_fresh_run(workspace):
    """reports/examples/ must be real tool output; regenerate with scripts/regenerate_examples.py."""
    for arguments, expected_exit in _load_script("regenerate_examples").COMMANDS:
        assert main(arguments) == expected_exit, arguments
    fresh_dir = workspace / "reports" / "examples"
    assert sorted(p.name for p in fresh_dir.iterdir()) == sorted(p.name for p in EXAMPLES.iterdir() if p.name != "README.md")
    for fresh in fresh_dir.iterdir():
        committed = EXAMPLES / fresh.name
        if fresh.suffix == ".json":
            assert _normalized_json(fresh) == _normalized_json(committed), fresh.name
        elif fresh.suffix == ".md":
            assert _normalized_markdown(fresh) == _normalized_markdown(committed), fresh.name
        elif fresh.suffix == ".csv":
            assert fresh.read_bytes() == committed.read_bytes(), fresh.name
        else:
            assert read_table(fresh).data.values.tolist() == read_table(committed).data.values.tolist(), fresh.name


def test_xlsx_generator_reproduces_the_committed_sample(tmp_path):
    _load_script("build_sample_xlsx").build(tmp_path / "orders.xlsx")
    for sheet in ("Read me", "Orders"):
        fresh = [[c.value for c in row] for row in load_workbook(tmp_path / "orders.xlsx")[sheet].iter_rows()]
        committed = [[c.value for c in row] for row in load_workbook(SAMPLE / "orders_raw.xlsx")[sheet].iter_rows()]
        assert fresh == committed


def test_sample_data_is_synthetic_and_uses_reserved_email_domains():
    emails = read_table(SAMPLE / "customers_raw.csv").data["Email"].dropna()
    domains = {e.strip().lower().split("@")[-1] for e in emails if e.strip()}
    assert domains <= {"example.com", "example"}


def test_cli_output_survives_a_legacy_console_encoding(tmp_path):
    """Regression: printing 'Kovač' to a cp1252 pipe crashed the profile command."""
    env = {**os.environ, "PYTHONIOENCODING": "cp1252", "PYTHONPATH": str(ROOT / "src")}
    completed = subprocess.run(
        [sys.executable, "-m", "data_cleaning", "profile", str(SAMPLE / "customers_raw.csv"), "--report-dir", str(tmp_path)],
        capture_output=True, env=env, cwd=tmp_path,
    )
    assert completed.returncode == 0, completed.stderr
    assert b"Kova\\u010d" in completed.stdout  # escaped on the console...
    assert "Kovač" in (tmp_path / "customers_raw_profile.md").read_text(encoding="utf-8")  # ...exact in the file
