"""Reading and writing CSV/XLSX files without silent type conversion.

Design rules:

* Every cell is loaded as text (pandas ``string`` dtype). No column is converted
  to a number or a date on import, so identifiers keep their leading zeros and
  nothing is guessed. Conversions only happen later, through explicit rules.
* Empty cells are loaded as empty strings; deciding what counts as *missing* is
  the job of :func:`data_cleaning.cleaning.standardize_missing`.
* Structurally malformed CSV files (rows with the wrong number of fields, broken
  quoting, duplicate header names) raise :class:`MalformedInputError` instead of
  being padded or truncated.
* Exports are written to a temporary file, reopened and compared with the data in
  memory, and only then moved to the requested path. Existing files are never
  overwritten unless explicitly allowed, and the input file is never a valid target.
"""

from __future__ import annotations

import codecs
import csv
import datetime as dt
import io
import logging
import os
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence
from xml.etree import ElementTree

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException

from .errors import (
    EmptyInputError,
    ExportError,
    InputFileError,
    MalformedInputError,
    OutputExistsError,
)
from .models import SOURCE_ROW

logger = logging.getLogger(__name__)

SUPPORTED_FORMATS = {".csv": "csv", ".xlsx": "xlsx"}

_UNSUPPORTED_HINTS = {
    ".xls": "Legacy .xls workbooks are not supported; save the file as .xlsx or .csv.",
    ".xlsm": "Macro-enabled workbooks (.xlsm) are not supported; save a copy as .xlsx.",
    ".xlsb": "Binary workbooks (.xlsb) are not supported; save a copy as .xlsx.",
    ".ods": "OpenDocument spreadsheets are not supported; save a copy as .xlsx or .csv.",
}

_CANDIDATE_DELIMITERS = ",;\t|"
_SNIFF_SAMPLE_CHARS = 64 * 1024
_EXCEL_ERROR_VALUES = {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"}
_MAX_LISTED_PROBLEMS = 10


@dataclass
class LoadedTable:
    """A dataset loaded as text plus metadata describing how it was read."""

    data: pd.DataFrame
    path: Path
    file_format: str
    encoding: str | None = None
    delimiter: str | None = None
    delimiter_detected: bool = False
    sheet: str | None = None
    available_sheets: list[str] = field(default_factory=list)
    skipped_blank_rows: int = 0
    formula_cells: int = 0
    uncalculated_formula_cells: int = 0
    warnings: list[str] = field(default_factory=list)

    def describe(self) -> dict[str, Any]:
        """Return JSON-friendly metadata about the read operation."""
        return {
            "path": self.path.as_posix(),
            "file_name": self.path.name,
            "format": self.file_format,
            "encoding": self.encoding,
            "delimiter": self.delimiter,
            "delimiter_detected": self.delimiter_detected,
            "sheet": self.sheet,
            "available_sheets": list(self.available_sheets),
            "rows": int(self.data.shape[0]),
            "columns": int(self.data.shape[1]),
            "skipped_blank_rows": self.skipped_blank_rows,
            "formula_cells": self.formula_cells,
            "uncalculated_formula_cells": self.uncalculated_formula_cells,
            "warnings": list(self.warnings),
        }


@dataclass
class ExportResult:
    """Outcome of a verified export."""

    path: Path
    file_format: str
    rows: int
    columns: int
    verified: bool
    warnings: list[str] = field(default_factory=list)

    def describe(self) -> dict[str, Any]:
        return {
            "path": self.path.as_posix(),
            "file_name": self.path.name,
            "format": self.file_format,
            "rows": self.rows,
            "columns": self.columns,
            "verified": self.verified,
            "warnings": list(self.warnings),
        }


# --------------------------------------------------------------------------- reading


def detect_format(path: str | os.PathLike[str]) -> str:
    """Return ``"csv"`` or ``"xlsx"`` from the file extension, or raise ``InputFileError``."""
    suffix = Path(path).suffix.lower()
    if suffix in SUPPORTED_FORMATS:
        return SUPPORTED_FORMATS[suffix]
    hint = _UNSUPPORTED_HINTS.get(suffix, "Supported formats are .csv and .xlsx.")
    raise InputFileError(f"Unsupported file type '{suffix or '(none)'}' for {Path(path).name}. {hint}")


def read_table(
    path: str | os.PathLike[str],
    *,
    sheet: str | None = None,
    encoding: str | None = None,
    delimiter: str | None = None,
    skip_blank_rows: bool = True,
) -> LoadedTable:
    """Read a CSV or XLSX file into a DataFrame of text values.

    Args:
        path: Input file path (``.csv`` or ``.xlsx``).
        sheet: Worksheet name for XLSX files. Defaults to the first sheet, with a
            warning when the workbook contains several sheets.
        encoding: CSV text encoding. When omitted, UTF-8 (with or without BOM) is
            tried first, then UTF-16 if a UTF-16 BOM is present, then Windows-1252
            with a warning.
        delimiter: CSV delimiter. When omitted it is detected among ``, ; TAB |``.
        skip_blank_rows: Skip rows in which every cell is empty (counted in
            ``LoadedTable.skipped_blank_rows``).

    Returns:
        A :class:`LoadedTable` whose ``data`` has one ``string`` column per header
        and an index named ``source_row`` holding the original row/line numbers.

    Raises:
        InputFileError: Path missing, not a file, unsupported type or unreadable.
        EmptyInputError: The file has no header or no data rows.
        MalformedInputError: The content is not a well-formed table.
    """
    file_path = Path(path)
    file_format = detect_format(file_path)
    _check_readable(file_path)
    if file_format == "csv":
        if sheet is not None:
            raise InputFileError("A sheet name can only be used with .xlsx files.")
        table = _read_csv(file_path, encoding=encoding, delimiter=delimiter, skip_blank_rows=skip_blank_rows)
    else:
        if encoding is not None or delimiter is not None:
            raise InputFileError("Encoding and delimiter options only apply to .csv files.")
        table = _read_xlsx(file_path, sheet=sheet, skip_blank_rows=skip_blank_rows)
    logger.info("Read %s: %d rows x %d columns", file_path.name, *table.data.shape)
    return table


def _check_readable(path: Path) -> None:
    if not path.exists():
        raise InputFileError(f"Input file not found: {path}")
    if not path.is_file():
        raise InputFileError(f"Input path is not a file: {path}")
    try:
        size = path.stat().st_size
        with path.open("rb"):
            pass
    except PermissionError as exc:
        raise InputFileError(f"Permission denied when reading {path}. Close the file if it is open in another program.") from exc
    except OSError as exc:
        raise InputFileError(f"Cannot open {path}: {exc}") from exc
    if size == 0:
        raise EmptyInputError(f"Input file is empty (0 bytes): {path.name}")


def _decode(raw: bytes, encoding: str | None, name: str) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    if encoding is not None:
        try:
            codecs.lookup(encoding)
        except LookupError as exc:
            raise InputFileError(f"Unknown encoding '{encoding}'.") from exc
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            raise MalformedInputError(
                f"{name} is not valid {encoding} (byte {exc.start}). Check the --encoding option."
            ) from exc
        return text.removeprefix("﻿"), encoding, warnings

    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        try:
            return raw.decode("utf-16").removeprefix("﻿"), "utf-16", warnings
        except UnicodeDecodeError as exc:
            raise MalformedInputError(f"{name} starts with a UTF-16 marker but is not valid UTF-16.") from exc
    try:
        return raw.decode("utf-8-sig"), "utf-8", warnings
    except UnicodeDecodeError:
        pass
    try:
        text = raw.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise MalformedInputError(
            f"Could not determine the text encoding of {name}. Specify it explicitly (e.g. --encoding latin-1)."
        ) from exc
    warnings.append(
        f"{name} is not valid UTF-8; it was decoded as Windows-1252 (cp1252). "
        "Check accented characters in the output, or specify the encoding explicitly."
    )
    return text, "cp1252", warnings


def _detect_delimiter(text: str) -> tuple[str, bool, list[str]]:
    sample = text[:_SNIFF_SAMPLE_CHARS]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=_CANDIDATE_DELIMITERS)
        return dialect.delimiter, True, []
    except csv.Error:
        # The sniffer fails on inconsistent rows. Fall back to the header line when it
        # contains exactly one candidate, so that row-level problems get precise errors.
        first_line = sample.splitlines()[0] if sample.splitlines() else ""
        present = [candidate for candidate in _CANDIDATE_DELIMITERS if candidate in first_line]
        if len(present) == 1:
            return present[0], True, [f"Delimiter {present[0]!r} was inferred from the header line only."]
        if present:
            raise MalformedInputError(
                "Could not reliably detect the CSV delimiter. Specify it explicitly (e.g. --delimiter ';')."
            ) from None
        return ",", False, ["No delimiter detected; the file was read as a single column."]


def _read_csv(path: Path, *, encoding: str | None, delimiter: str | None, skip_blank_rows: bool) -> LoadedTable:
    text, encoding_used, warnings = _decode(path.read_bytes(), encoding, path.name)
    if "\x00" in text:
        raise MalformedInputError(
            f"{path.name} contains NUL bytes; it may be binary or use another encoding (e.g. utf-16)."
        )
    if not text.strip():
        raise EmptyInputError(f"Input file contains no data: {path.name}")

    if delimiter is not None:
        if len(delimiter) != 1:
            raise InputFileError(f"The delimiter must be a single character, got {delimiter!r}.")
        delimiter_used, detected = delimiter, False
    else:
        delimiter_used, detected, delimiter_warnings = _detect_delimiter(text)
        warnings.extend(delimiter_warnings)

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter_used, strict=True)
    records: list[tuple[int, list[str]]] = []
    previous_line = 0
    try:
        for record in reader:
            start_line = previous_line + 1
            previous_line = reader.line_num
            if not record:  # a physically empty line
                continue
            records.append((start_line, record))
    except csv.Error as exc:
        raise MalformedInputError(f"{path.name}: invalid CSV near line {reader.line_num}: {exc}") from exc

    if not records:
        raise EmptyInputError(f"Input file contains no data: {path.name}")
    header_line, header = records[0]
    bad_rows = [(line, len(fields)) for line, fields in records[1:] if len(fields) != len(header)]
    if bad_rows:
        listed = ", ".join(f"line {line} ({count} fields)" for line, count in bad_rows[:_MAX_LISTED_PROBLEMS])
        more = f" and {len(bad_rows) - _MAX_LISTED_PROBLEMS} more" if len(bad_rows) > _MAX_LISTED_PROBLEMS else ""
        raise MalformedInputError(
            f"{path.name}: the header has {len(header)} fields but {len(bad_rows)} row(s) do not: "
            f"{listed}{more}. Check the delimiter and quoting."
        )
    data, skipped = _build_frame(header, records[1:], warnings, path.name, skip_blank_rows)
    return LoadedTable(
        data=data,
        path=path,
        file_format="csv",
        encoding=encoding_used,
        delimiter=delimiter_used,
        delimiter_detected=detected,
        skipped_blank_rows=skipped,
        warnings=warnings,
    )


def _read_xlsx(path: Path, *, sheet: str | None, skip_blank_rows: bool) -> LoadedTable:
    warnings: list[str] = []
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise MalformedInputError(
            f"{path.name} could not be opened as an .xlsx workbook. The file may be corrupt, "
            f"password-protected, or not a real .xlsx file ({type(exc).__name__})."
        ) from exc
    try:
        sheet_names = list(workbook.sheetnames)
        if sheet is None:
            sheet = sheet_names[0]
            if len(sheet_names) > 1:
                warnings.append(
                    f"The workbook has {len(sheet_names)} sheets; only the first ('{sheet}') was read. "
                    "Use the sheet option to choose another."
                )
        elif sheet not in sheet_names:
            raise InputFileError(f"Sheet '{sheet}' not found in {path.name}. Available sheets: {', '.join(sheet_names)}")
        worksheet = workbook[sheet]
        if not hasattr(worksheet, "iter_rows"):
            raise InputFileError(f"Sheet '{sheet}' is a chart sheet and contains no table.")
        records: list[tuple[int, list[str]]] = []
        excel_errors = 0
        for row_number, values in enumerate(worksheet.iter_rows(min_row=1, values_only=True), start=1):
            texts = [_excel_value_to_text(value) for value in values]
            excel_errors += sum(1 for value in texts if value in _EXCEL_ERROR_VALUES)
            records.append((row_number, texts))
    finally:
        workbook.close()

    formulas, uncalculated = _count_formula_cells(path, sheet)
    if formulas is None:
        warnings.append(f"Formula cells in sheet '{sheet}' could not be inspected; formulas, if any, were read as their last saved values.")
    elif formulas:
        warnings.append(f"{formulas} cell(s) in sheet '{sheet}' contain formulas; their last values saved by Excel were used.")
    if uncalculated:
        warnings.append(
            f"{uncalculated} formula cell(s) in sheet '{sheet}' have no saved value (the workbook was probably created by a "
            "program and never recalculated in Excel); they were read as EMPTY. Open and save the file in Excel first, "
            "or treat these values as unknown."
        )

    first = next((index for index, (_, fields) in enumerate(records) if any(fields)), None)
    if first is None:
        raise EmptyInputError(f"Sheet '{sheet}' in {path.name} contains no data.")
    header_row, header = records[first]
    if header_row != 1:
        warnings.append(f"The header was found on row {header_row} of sheet '{sheet}' (rows above it are empty).")
    body = records[first + 1 :]
    # Drop trailing columns that have neither a header nor any value (formatting artefacts).
    width = len(header)
    while width > 1 and header[width - 1] == "" and all(fields[width - 1] == "" for _, fields in body):
        width -= 1
    header = header[:width]
    body = [(number, fields[:width]) for number, fields in body]
    if excel_errors:
        warnings.append(f"{excel_errors} cell(s) contain Excel error values such as #N/A or #DIV/0!.")
    data, skipped = _build_frame(header, body, warnings, path.name, skip_blank_rows)
    return LoadedTable(
        data=data,
        path=path,
        file_format="xlsx",
        sheet=sheet,
        available_sheets=sheet_names,
        skipped_blank_rows=skipped,
        formula_cells=formulas or 0,
        uncalculated_formula_cells=uncalculated or 0,
        warnings=warnings,
    )


_SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _count_formula_cells(path: Path, sheet: str) -> tuple[int | None, int | None]:
    """Count formula cells in ``sheet``, and those without a saved (cached) value.

    openpyxl, used with ``data_only=True``, returns the value Excel saved for each
    formula and ``None`` when there is none, which is indistinguishable from an empty
    cell. This scan of the sheet XML makes that case visible. Returns ``(None, None)``
    if the workbook structure cannot be inspected.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {rel.get("Id"): rel.get("Target", "") for rel in rels}
            rel_id = next(s.get(f"{_REL_NS}id") for s in workbook.iter(f"{_SHEET_NS}sheet") if s.get("name") == sheet)
            target = targets[rel_id]
            member = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            formulas = uncalculated = 0
            with archive.open(member) as handle:
                for _, element in ElementTree.iterparse(handle):
                    if element.tag == f"{_SHEET_NS}c":
                        if element.find(f"{_SHEET_NS}f") is not None:
                            formulas += 1
                            value = element.find(f"{_SHEET_NS}v")
                            if value is None or not value.text:
                                uncalculated += 1
                        element.clear()
                    elif element.tag == f"{_SHEET_NS}row":
                        element.clear()  # release the row's (already cleared) cells: keeps memory flat
            return formulas, uncalculated
    except (KeyError, StopIteration, zipfile.BadZipFile, ElementTree.ParseError, OSError):
        return None, None


def _build_frame(
    header: list[str],
    body: Iterable[tuple[int, list[str]]],
    warnings: list[str],
    name: str,
    skip_blank_rows: bool,
) -> tuple[pd.DataFrame, int]:
    columns = list(header)
    for position, column in enumerate(columns, start=1):
        if column.strip() == "":
            columns[position - 1] = f"column_{position}"
            warnings.append(f"Column {position} has no header; it was named 'column_{position}'.")
    duplicates = sorted({column for column in columns if columns.count(column) > 1})
    if duplicates:
        raise MalformedInputError(
            f"{name} has duplicate column names: {', '.join(repr(c) for c in duplicates)}. "
            "Rename them in the source file so that every column is uniquely identified."
        )
    rows: list[list[str]] = []
    row_numbers: list[int] = []
    skipped = 0
    for number, fields in body:
        if skip_blank_rows and all(value == "" for value in fields):
            skipped += 1
            continue
        rows.append(fields)
        row_numbers.append(number)
    if skipped:
        warnings.append(f"{skipped} completely empty row(s) were skipped.")
    if not rows:
        raise EmptyInputError(f"{name} has a header but no data rows.")
    frame = pd.DataFrame(rows, columns=columns, dtype="string")
    frame.index = pd.Index(row_numbers, name=SOURCE_ROW)
    return frame, skipped


def _excel_value_to_text(value: Any) -> str:
    """Convert a cell value from openpyxl to text, without losing information."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format_float(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def format_float(value: float) -> str:
    """Canonical text for a float: integral values without ``.0``, others with the shortest exact repr."""
    if value != value:  # NaN
        return ""
    if value.is_integer() and abs(value) < 1e16:
        return str(int(value))
    return repr(value)


# --------------------------------------------------------------------------- writing


def render_text(data: pd.DataFrame) -> pd.DataFrame:
    """Return the canonical text representation of ``data`` (missing values as ``<NA>``).

    This is the representation written to CSV and used to verify exported files.
    """
    rendered = {}
    for column in data.columns:
        series = data[column]
        if pd.api.types.is_float_dtype(series.dtype):
            rendered[column] = series.map(lambda v: pd.NA if pd.isna(v) else format_float(float(v))).astype("string")
        elif pd.api.types.is_bool_dtype(series.dtype):
            rendered[column] = series.map(lambda v: pd.NA if pd.isna(v) else ("TRUE" if v else "FALSE")).astype("string")
        else:
            rendered[column] = series.astype("string")
    return pd.DataFrame(rendered, index=data.index)


def write_table(
    data: pd.DataFrame,
    path: str | os.PathLike[str],
    *,
    overwrite: bool = False,
    protected_paths: Sequence[str | os.PathLike[str]] = (),
    csv_encoding: str = "utf-8",
    csv_delimiter: str = ",",
    sheet_name: str = "cleaned",
) -> ExportResult:
    """Write ``data`` to CSV or XLSX, verify the written file, then move it into place.

    Args:
        data: The DataFrame to export (its index is not written).
        path: Destination path; the extension selects the format.
        overwrite: Allow replacing an existing file. Never applies to ``protected_paths``.
        protected_paths: Paths that must never be written to (typically the input file).
        csv_encoding: Encoding for CSV output (``utf-8-sig`` helps Excel detect UTF-8).
        csv_delimiter: Delimiter for CSV output.
        sheet_name: Worksheet name for XLSX output.

    Raises:
        OutputExistsError: ``path`` exists and ``overwrite`` is False.
        ExportError: ``path`` is protected, cannot be written, or fails verification.
    """
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise ExportError(f"Unsupported output type '{suffix or '(none)'}'. Use .csv or .xlsx.")
    file_format = SUPPORTED_FORMATS[suffix]
    check_output_path(target, overwrite=overwrite, protected_paths=protected_paths)
    if len(csv_delimiter) != 1:
        raise ExportError(f"The CSV delimiter must be a single character, got {csv_delimiter!r}.")

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportError(f"Cannot create output directory {target.parent}: {exc}") from exc

    temporary = target.with_name(f".{target.stem}.tmp-{uuid.uuid4().hex[:8]}{target.suffix}")
    try:
        if file_format == "csv":
            _write_csv(data, temporary, encoding=csv_encoding, delimiter=csv_delimiter)
        else:
            _write_xlsx(data, temporary, sheet_name=sheet_name)
        warnings = _verify_export(data, temporary, file_format, csv_encoding, csv_delimiter, sheet_name)
        os.replace(temporary, target)
    except PermissionError as exc:
        raise ExportError(f"Permission denied when writing {target}. Close the file if it is open in another program.") from exc
    except OSError as exc:
        raise ExportError(f"Could not write {target}: {exc}") from exc
    finally:
        if temporary.exists():
            temporary.unlink()
    logger.info("Wrote and verified %s (%d rows x %d columns)", target.name, *data.shape)
    return ExportResult(target, file_format, int(data.shape[0]), int(data.shape[1]), True, warnings)


def check_output_path(
    path: str | os.PathLike[str],
    *,
    overwrite: bool = False,
    protected_paths: Sequence[str | os.PathLike[str]] = (),
) -> None:
    """Raise if ``path`` is protected, is a directory, or exists without ``overwrite``.

    Used to check every planned output *before* any processing starts, so that a
    run never stops halfway with some files written and others refused.
    """
    target = Path(path)
    for protected in protected_paths:
        if _same_path(target, Path(protected)):
            raise ExportError(f"Refusing to write to {target}: it is the input file, which is never modified.")
    if target.exists():
        if target.is_dir():
            raise ExportError(f"Output path is a directory: {target}")
        if not overwrite:
            raise OutputExistsError(
                f"Output file already exists: {target}. Choose another name or allow overwriting explicitly (--overwrite)."
            )


def write_text_file(
    path: str | os.PathLike[str],
    content: str,
    *,
    overwrite: bool = False,
    protected_paths: Sequence[str | os.PathLike[str]] = (),
) -> Path:
    """Atomically write a UTF-8 text file (used for reports) with overwrite protection."""
    target = Path(path)
    check_output_path(target, overwrite=overwrite, protected_paths=protected_paths)
    temporary = target.with_name(f".{target.stem}.tmp-{uuid.uuid4().hex[:8]}{target.suffix}")
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary.write_text(content, encoding="utf-8", newline="\n")
        os.replace(temporary, target)
    except PermissionError as exc:
        raise ExportError(f"Permission denied when writing {target}. Close the file if it is open in another program.") from exc
    except OSError as exc:
        raise ExportError(f"Could not write {target}: {exc}") from exc
    finally:
        if temporary.exists():
            temporary.unlink()
    return target


def _same_path(a: Path, b: Path) -> bool:
    try:
        if a.exists() and b.exists():
            return os.path.samefile(a, b)
        return os.path.normcase(a.resolve()) == os.path.normcase(b.resolve())
    except OSError:
        return False


def _write_csv(data: pd.DataFrame, path: Path, *, encoding: str, delimiter: str) -> None:
    text = render_text(data)
    try:
        with path.open("w", encoding=encoding, newline="") as handle:
            writer = csv.writer(handle, delimiter=delimiter, lineterminator="\n")
            writer.writerow([str(column) for column in text.columns])
            for row in text.itertuples(index=False, name=None):
                writer.writerow(["" if pd.isna(value) else value for value in row])
    except (UnicodeEncodeError, LookupError) as exc:
        raise ExportError(f"Cannot encode the output as {encoding}: {exc}") from exc


def _write_xlsx(data: pd.DataFrame, path: Path, *, sheet_name: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    columns = list(data.columns)
    numeric = {c for c in columns if pd.api.types.is_numeric_dtype(data[c].dtype) and not pd.api.types.is_bool_dtype(data[c].dtype)}
    text = render_text(data)
    for col_index, column in enumerate(columns, start=1):
        _set_text_cell(worksheet, 1, col_index, str(column))
    for row_index, position in enumerate(range(len(data)), start=2):
        for col_index, column in enumerate(columns, start=1):
            value = data[column].iat[position]
            if pd.isna(value):
                continue
            if column in numeric:
                number = value.item() if hasattr(value, "item") else value
                worksheet.cell(row=row_index, column=col_index, value=number)
            else:
                try:
                    _set_text_cell(worksheet, row_index, col_index, text[column].iat[position])
                except IllegalCharacterError as exc:
                    raise ExportError(
                        f"Column '{column}', data row {position + 1} contains a control character that "
                        "cannot be stored in .xlsx. Export to CSV or clean that value first."
                    ) from exc
    workbook.save(path)


def _set_text_cell(worksheet: Any, row: int, column: int, value: str) -> None:
    cell = worksheet.cell(row=row, column=column, value=value)
    # Text such as "=SUM(A1)" must stay text, not become a live formula.
    cell.data_type = "s"


def _verify_export(
    data: pd.DataFrame, path: Path, file_format: str, encoding: str, delimiter: str, sheet_name: str
) -> list[str]:
    """Reopen the written file and compare it cell by cell with ``data``."""
    warnings: list[str] = []
    expected = render_text(data).reset_index(drop=True).fillna("")
    blank_rows = expected.eq("").all(axis=1)
    if blank_rows.any():
        warnings.append(f"{int(blank_rows.sum())} exported row(s) are entirely empty.")
        expected = expected[~blank_rows].reset_index(drop=True)
    try:
        if file_format == "csv":
            reloaded = read_table(path, encoding=encoding, delimiter=delimiter).data
        else:
            reloaded = read_table(path, sheet=sheet_name).data
    except EmptyInputError:
        if expected.empty:
            return warnings
        raise ExportError(f"Verification failed: {path.name} could not be read back as a table.") from None
    except InputFileError as exc:
        raise ExportError(f"Verification failed: the written file could not be reopened ({exc}).") from exc
    reloaded = reloaded.reset_index(drop=True).fillna("")
    expected_columns = [str(c) for c in data.columns]
    if list(reloaded.columns) != expected_columns:
        raise ExportError(
            f"Verification failed: columns read back {list(reloaded.columns)} differ from {expected_columns}."
        )
    if reloaded.shape != expected.shape:
        raise ExportError(f"Verification failed: read back {reloaded.shape[0]} rows, expected {expected.shape[0]}.")
    expected.columns = expected_columns
    mismatch = expected.astype(str).ne(reloaded.astype(str))
    if mismatch.to_numpy().any():
        row, col = next(zip(*mismatch.to_numpy().nonzero()))
        column = expected_columns[col]
        raise ExportError(
            f"Verification failed: column '{column}', data row {row + 1}: wrote {expected.iat[row, col]!r} "
            f"but read back {reloaded.iat[row, col]!r}."
        )
    return warnings
