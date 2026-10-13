"""Generate the synthetic benchmark datasets and a manifest of expected outcomes.

Deterministic (fixed seed): the same arguments always produce the same rows.
The schemas are those of data/sample/customers_raw.csv and of the 'Orders' sheet of
data/sample/orders_raw.xlsx, so config/customers_config.json and config/orders_config.json
apply unchanged. Each row's defects are chosen here and the expected counts in
manifest.json are derived from those choices, not from the toolkit.

Usage:
    python scripts/benchmark/generate_data.py WORKDIR [--rows 100000]

Writes WORKDIR/data/customers_100k.csv (or customers_<rows>.csv), orders_<...>.xlsx and
manifest.json. Keep WORKDIR outside the repository: the files are large.
"""
import argparse
import csv
import datetime as dt
import hashlib
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

SEED = 20261013
parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
parser.add_argument("workdir", type=Path, help="benchmark working directory (outside the repository)")
parser.add_argument("--rows", type=int, default=100_000, help="data rows per dataset (default 100000)")
cli_args = parser.parse_args()
if cli_args.rows < 2_000:
    parser.error("--rows must be at least 2000")
TOTAL_ROWS = cli_args.rows
EXACT_DUP_COPIES = TOTAL_ROWS // 100
SUFFIX = "100k" if TOTAL_ROWS == 100_000 else str(TOTAL_ROWS)
out = cli_args.workdir / "data"
out.mkdir(parents=True, exist_ok=True)

FIRST = ["Ana", "Ben", "Chloé", "Dev", "Elif", "Femi", "Grace", "Hana", "Ivan", "Jo", "Kai", "Lena", "Mo", "Nia", "Omar",
         "Pia", "Quinn", "Rui", "Sara", "Tomás", "Uma", "Vik", "Wen", "Xavi", "Yara", "Zoé"]
LAST = ["Kovač", "Okafor", "Martin", "Patel", "Demir", "Adeyemi", "Liu", "Sato", "Petrov", "Kim", "Müller", "Novak",
        "Haddad", "Silva", "Rossi", "Nguyen", "García", "Dubois", "Jensen", "Moreau"]
CITY_VARIANTS = {"New York": ["New York", " new york", "NEW YORK", "NYC", "nyc "], "Boston": ["Boston", "boston"],
                 "Chicago": ["Chicago", "CHICAGO"], "Seattle": ["Seattle", "seattle"], "Austin": ["Austin"], "Denver": ["Denver", " Denver "]}
STATUS_VARIANTS = ["active", "Active", "ACTIVE", "inactive", "Inactive", "closed", " closed"]


def customers(rng: random.Random) -> dict:
    base_count = TOTAL_ROWS - EXACT_DUP_COPIES
    rows, flags = [], []
    clean_indexes = []
    keydup_sources_used = set()
    keydup_target = max(1, TOTAL_ROWS * 500 // 100_000)
    for i in range(base_count):
        f = set()
        r = rng.random
        # identifier: sequential, zero-padded (leading zeros for all ids < 100000)
        cid = f"{i + 1:06d}"
        if r() < 0.005:
            cid, _ = "", f.add("missing_id")
        first, last = rng.choice(FIRST), rng.choice(LAST)
        name = f"{first} {last}" if r() > 0.05 else f"  {first.lower()}   {last} "
        email = f"{first.lower()}.{last.lower()}{i}@example.com"
        x = r()
        if x < 0.005:
            email = ""; f.add("missing_email")
        elif x < 0.010:
            email = f"{first.lower()}{i}@example"; f.add("bad_email")
        elif x < 0.110:
            email = email.upper()
        city_key = rng.choice(list(CITY_VARIANTS))
        city = rng.choice(CITY_VARIANTS[city_key]) if r() > 0.005 else "Bostn"
        if city == "Bostn":
            f.add("unmapped_city")
        # signup date
        x = r()
        year, month = rng.randint(2016, 2025), rng.randint(1, 12)
        if x < 0.50:
            date = f"{year}-{month:02d}-{rng.randint(1, 28):02d}"
        elif x < 0.65:
            date = f"{rng.randint(13, 28):02d}/{month:02d}/{year}"  # only DD/MM fits
        elif x < 0.80:
            date = f"{month:02d}/{rng.randint(13, 28):02d}/{year}"  # only MM/DD fits
        elif x < 0.85:
            a = rng.randint(1, 12); b = rng.choice([v for v in range(1, 13) if v != a])
            date = f"{a:02d}/{b:02d}/{year}"; f.add("ambiguous_date")
        elif x < 0.86:
            date = f"{month:02d}/{month:02d}/{year}"  # same day and month: both readings agree
        elif x < 0.88:
            date = rng.choice(["2024-02-30", "2023-13-01", "yesterday", "31/02/2022"]); f.add("invalid_date")
        elif x < 0.89:
            date = f"{rng.randint(2031, 2035)}-{month:02d}-15"; f.add("future_date")
        else:
            date = f"{year}-{month:02d}-{rng.randint(1, 28):02d}"
        # annual spend
        x = r()
        amount = rng.randint(1, 9999)
        cents = rng.randint(0, 99)
        if x < 0.40:
            spend = str(amount)
        elif x < 0.60:
            spend = f"{amount:,}.{cents:02d}"
        elif x < 0.70:
            spend = f"${amount}"
        elif x < 0.76:
            spend = rng.choice(["N/A", "NULL", "   ", "n/a", ""]); f.add("missing_spend")
        elif x < 0.77:
            spend = f"{amount:,}.{cents:02d}".replace(",", "X").replace(".", ",").replace("X", "."); f.add("eu_spend")
        elif x < 0.78:
            spend = rng.choice(["twelve hundred", "approx 500", "tbd"]); f.add("word_spend")
        elif x < 0.79:
            spend = f"-{amount}"; f.add("negative_spend")
        elif x < 0.80:
            spend = "0"
        else:
            spend = f"{amount}.{cents:02d}"
        status = rng.choice(STATUS_VARIANTS)
        if "missing_id" not in f and r() < 0.01:
            status = rng.choice(["on hold", "unknown"]); f.add("bad_status")
        postal = f"{rng.randint(0, 99999):05d}"
        if r() < 0.005:
            postal = str(rng.randint(1000, 9999)); f.add("short_postal")
        # key duplicate: reuse the id of an earlier clean row, with a different amount
        if not f and i > 1000 and len(keydup_sources_used) < keydup_target and r() < 0.006:
            source = rng.choice(clean_indexes)
            if source not in keydup_sources_used:
                keydup_sources_used.add(source)
                cid = rows[source][0]
                spend = str(int(rows[source][5].replace(",", "").replace("$", "").split(".")[0] or 0) + 1) if rows[source][5].strip() else "1"
                f.add("key_dup")
                flags[source].add("key_dup_source")
        rows.append([cid, name, email, city, date, spend, status, postal])
        flags.append(f)
        if not f:
            clean_indexes.append(i)
    # exact duplicate copies of clean rows that are not involved in key duplicates
    candidates = [i for i in clean_indexes if not flags[i]]
    sources = sorted(rng.sample(candidates, EXACT_DUP_COPIES))
    final_rows, final_flags = [], []
    source_set = set(sources)
    for i, (row, f) in enumerate(zip(rows, flags)):
        final_rows.append(row); final_flags.append(f)
        if i in source_set:
            final_rows.append(list(row)); final_flags.append({"exact_dup_copy"})
    header = ["Customer ID", " Full Name ", "Email", "City", "Signup Date", "Annual  Spend", "Status", "Postal Code"]
    path = out / f"customers_{SUFFIX}.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(final_rows)

    def count(flag):
        return sum(flag in f for f in final_flags)

    present = [f for f in final_flags if "exact_dup_copy" not in f]  # rows present at validation
    validation = {
        "required_columns": 0,
        "id_present": count("missing_id"),
        "id_unique": 2 * count("key_dup"),
        "email_present": count("missing_email"),
        "email_format": count("bad_email"),
        "status_allowed": count("bad_status"),
        "signup_is_iso_date": count("ambiguous_date") + count("invalid_date"),
        "signup_in_range": count("ambiguous_date") + count("invalid_date") + count("future_date"),
        "spend_non_negative": count("eu_spend") + count("word_spend") + count("negative_spend"),
        "postal_code_5_digits": count("short_postal"),
    }
    issue_flags = {"missing_id", "key_dup", "key_dup_source", "missing_email", "bad_email", "bad_status", "ambiguous_date",
                   "invalid_date", "future_date", "eu_spend", "word_spend", "negative_spend", "short_postal"}
    quarantined = count("missing_id") + count("bad_status")
    return {
        "file": path.name,
        "rows": len(final_rows),
        "columns": len(header),
        "row_flags": [sorted(f) for f in final_flags],
        "expected": {
            "removed": EXACT_DUP_COPIES,
            "quarantined": quarantined,
            "cleaned": len(final_rows) - EXACT_DUP_COPIES - quarantined,
            "key_duplicate_groups": count("key_dup"),
            "date_failures": count("ambiguous_date") + count("invalid_date"),
            "ambiguous_dates": count("ambiguous_date"),
            "number_failures": count("eu_spend") + count("word_spend"),
            "validation": validation,
            "validation_issue_count": sum(validation.values()),
            "rows_with_issues": sum(bool(f & issue_flags) for f in present),
            "missing_spend_before": count("missing_spend"),
        },
    }


def orders(rng: random.Random) -> dict:
    base_count = TOTAL_ROWS - EXACT_DUP_COPIES
    rows, flags = [], []
    order_no, lines_left, skus_used = 0, 0, set()
    clean_indexes, keydup_target, keydups, keydup_sources_used = [], max(1, TOTAL_ROWS * 300 // 100_000), 0, set()
    for i in range(base_count):
        f = set()
        r = rng.random
        if lines_left == 0:
            order_no += 1
            lines_left = rng.randint(1, 3)
            skus_used = set()
            order_customer = rng.randint(1, 99999)
        lines_left -= 1
        oid = f"SO-{order_no:06d}"
        if r() < 0.005:
            oid = None; f.add("missing_order_id")
        x = r()
        day = dt.date(2024, 1, 1) + dt.timedelta(days=rng.randint(0, 365))
        if x < 0.60:
            date = dt.datetime(day.year, day.month, day.day)
        elif x < 0.85:
            date = day.strftime("%d.%m.%Y")
        elif x < 0.86:
            date = "2024-13-01"; f.add("invalid_date")
        elif x < 0.88:
            date = dt.datetime(day.year, day.month, day.day, rng.randint(8, 18), 30); f.add("datetime_with_time")
        else:
            date = day.isoformat()
        if r() < 0.03:
            customer = order_customer % 100000 or 1; f.add("numeric_customer_id")  # Excel number: zeros lost
        else:
            customer = f"{order_customer:06d}"
        sku_number = rng.choice([n for n in range(1, 400) if n not in skus_used])
        skus_used.add(sku_number)
        sku = f"AB-{sku_number:03d}"
        if r() < 0.10:
            sku = f" {sku.lower()}"
        x = r()
        if x < 0.85:
            qty = rng.randint(1, 20)
        elif x < 0.97:
            qty = str(rng.randint(1, 20))
        elif x < 0.98:
            qty = "two"; f.add("word_qty")
        else:
            qty = 0; f.add("zero_qty")
        x = r()
        if x < 0.85:
            price = round(rng.uniform(1, 500), 2)
        elif x < 0.98:
            price = f"{rng.randint(1, 500)}.{rng.randint(0, 99):02d}"
        else:
            price = None
        channel = rng.choice(["Online", "online", "Store", "STORE ", "Phone", "phone"])
        if "missing_order_id" not in f and r() < 0.005:
            channel = "Fax"; f.add("fax")
        if not f and oid and i > 1000 and keydups < keydup_target and r() < 0.004:
            source = rng.choice(clean_indexes)
            if source not in keydup_sources_used:  # one key duplicate per source -> groups of exactly 2
                keydup_sources_used.add(source)
                src = rows[source]
                oid, sku = src[0], src[3]
                qty = 21  # different quantity, so the rows are not exact duplicates
                f.add("key_dup"); keydups += 1
                flags[source].add("key_dup_source")
        rows.append([oid, date, customer, sku, qty, price, channel])
        flags.append(f)
        if not f:
            clean_indexes.append(i)
    candidates = [i for i in clean_indexes if not flags[i]]
    sources = set(rng.sample(candidates, EXACT_DUP_COPIES))
    final_rows, final_flags = [], []
    for i, (row, f) in enumerate(zip(rows, flags)):
        final_rows.append(row); final_flags.append(f)
        if i in sources:
            final_rows.append(list(row)); final_flags.append({"exact_dup_copy"})
    header = ["Order ID", "Order Date", "Customer ID", "SKU", "Quantity", "Unit Price", "Channel"]
    workbook = Workbook(write_only=True)
    notes = workbook.create_sheet("Read me")
    notes.append(["Synthetic benchmark data (seed %d). Data on sheet 'Orders'." % SEED])
    sheet = workbook.create_sheet("Orders")
    sheet.append(header)
    for row in final_rows:
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, dt.datetime):
                cell.number_format = "yyyy-mm-dd hh:mm" if value.hour else "yyyy-mm-dd"
            cells.append(cell)
        sheet.append(cells)
    path = out / f"orders_{SUFFIX}.xlsx"
    workbook.save(path)

    def count(flag):
        return sum(flag in f for f in final_flags)

    present = [f for f in final_flags if "exact_dup_copy" not in f]
    validation = {
        "order_id_present": count("missing_order_id"),
        "customer_id_6_digits": count("numeric_customer_id"),
        "quantity_at_least_1": count("word_qty") + count("zero_qty"),
        "channel_allowed": count("fax"),
        "order_date_is_iso_date": count("invalid_date") + count("datetime_with_time"),
    }
    issue_flags = {"missing_order_id", "numeric_customer_id", "word_qty", "zero_qty", "fax", "invalid_date", "datetime_with_time"}
    quarantined = sum(bool(f & {"missing_order_id", "fax"}) for f in final_flags)
    return {
        "file": path.name,
        "rows": len(final_rows),
        "columns": len(header),
        "row_flags": [sorted(f) for f in final_flags],
        "expected": {
            "removed": EXACT_DUP_COPIES,
            "quarantined": quarantined,
            "cleaned": len(final_rows) - EXACT_DUP_COPIES - quarantined,
            "key_duplicate_groups": count("key_dup"),
            "date_failures": count("invalid_date") + count("datetime_with_time"),
            "number_failures": count("word_qty"),
            "validation": validation,
            "validation_issue_count": sum(validation.values()),
            "rows_with_issues": sum(bool(f & issue_flags) for f in present),
        },
    }


manifest = {"seed": SEED, "customers": customers(random.Random(SEED)), "orders": orders(random.Random(SEED + 1))}
for name in ("customers", "orders"):
    manifest[name]["sha256"] = hashlib.sha256((out / manifest[name]["file"]).read_bytes()).hexdigest()
(out / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
for name in ("customers", "orders"):
    m = manifest[name]
    print(name, m["file"], m["rows"], "rows x", m["columns"], "cols;", json.dumps(m["expected"]))
