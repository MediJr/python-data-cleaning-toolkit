"""Independent verification of benchmark outputs. Deliberately does NOT import data_cleaning.

Usage:
    python scripts/benchmark/verify_outputs.py WORKDIR WORKFLOW RUN_NUMBER

WORKFLOW is "customers" or "orders". Expected values come from the generator's manifest and
from an independent re-derivation of each configured rule applied to the raw input; files are
read with csv/openpyxl only. Exit code 0 when every check passes, 1 otherwise.
"""

import csv
import datetime as dt
import hashlib
import json
import re
import sys
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

if len(sys.argv) != 4:
    sys.exit(__doc__)
WORKDIR = Path(sys.argv[1]).resolve()
workflow, run = sys.argv[2], int(sys.argv[3])
manifest = json.loads((WORKDIR / "data" / "manifest.json").read_text(encoding="utf-8"))[workflow]
expected = manifest["expected"]
flags = [set(f) for f in manifest["row_flags"]]
run_dir = WORKDIR / "runs" / workflow / f"run{run}"
checks, failures = 0, []


def check(ok, label):
    global checks
    checks += 1
    if not ok:
        failures.append(label)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fmt_float(v):
    return str(int(v)) if float(v).is_integer() else repr(float(v))


def excel_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, float):
        return fmt_float(v)
    return str(v)


def parse_date(text, formats):
    results = set()
    for f in formats:
        try:
            results.add(dt.datetime.strptime(text.strip(), f))
        except ValueError:
            pass
    return next(iter(results)).strftime("%Y-%m-%d") if len(results) == 1 else None


# ---- source preservation and statuses
source = WORKDIR / "data" / manifest["file"]
check(sha(source) == manifest["sha256"], "input SHA-256 unchanged since generation")
report = json.loads(next(run_dir.glob("*_report.json")).read_text(encoding="utf-8"))
check(report["input"]["sha256_before"] == report["input"]["sha256_after"] == manifest["sha256"], "report fingerprints")
check(report["input"]["unchanged"] is True, "report unchanged flag")
check(report["execution"]["status"] == "success", "execution success")
check(report["data_quality"]["status"] == "issues_found", "quality status")
results = json.loads((WORKDIR / "results.json").read_text(encoding="utf-8"))
check(next(r for r in results if r["workflow"] == workflow and r["run"] == run)["exit_code"] == 3, "exit code 3")

# ---- read input and outputs without the toolkit
if workflow == "customers":
    def read(path):
        with open(path, encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        return rows[0], rows[1:]

    raw_header, raw = read(source)
    header, cleaned = read(run_dir / "customers_clean.csv")
    q_header, quarantined = read(run_dir / "customers_clean_quarantine.csv")
    r_header, removed = read(run_dir / "customers_clean_removed_rows.csv")
    cols = ["customer_id", "full_name", "email", "city", "signup_date", "annual_spend", "status", "postal_code"]
    quarantine_flags = {"missing_id", "bad_status"}
else:
    def read_xlsx(path, sheet, with_types=False):
        worksheet = load_workbook(path, read_only=True)[sheet]
        return [[(c.value, c.data_type) if with_types else c.value for c in row] for row in worksheet.iter_rows()]

    raw_all = read_xlsx(source, "Orders")
    raw_header, raw = raw_all[0], raw_all[1:]
    out_all = read_xlsx(run_dir / "orders_clean.xlsx", "cleaned", with_types=True)
    header, cleaned_typed = [v for v, _ in out_all[0]], out_all[1:]
    cleaned = [[excel_text(v) for v, _ in row] for row in cleaned_typed]
    q_all = read_xlsx(run_dir / "orders_clean_quarantine.xlsx", "cleaned")
    q_header, quarantined = q_all[0], [[excel_text(v) for v in r] for r in q_all[1:]]
    r_all = read_xlsx(run_dir / "orders_clean_removed_rows.xlsx", "cleaned")
    r_header, removed = r_all[0], [[excel_text(v) for v in r] for r in r_all[1:]]
    cols = ["order_id", "order_date", "customer_id", "sku", "quantity", "unit_price", "channel"]
    quarantine_flags = {"missing_order_id", "fax"}

# ---- shape and row accounting
check(len(raw) == manifest["rows"], "input rows")
check(header == cols, f"cleaned header {header}")
check(list(q_header) == cols + ["quarantine_reason"], "quarantine header")
check(list(r_header) == cols + ["removal_reason"], "removed header")
check(len(cleaned) + len(removed) + len(quarantined) == len(raw), "row conservation")
check(
    (len(cleaned), len(removed), len(quarantined)) == (expected["cleaned"], expected["removed"], expected["quarantined"]),
    f"counts {(len(cleaned), len(removed), len(quarantined))} vs {(expected['cleaned'], expected['removed'], expected['quarantined'])}",
)
dims = report["dimensions"]
check(dims["initial"] == {"rows": len(raw), "columns": len(raw_header)}, "report initial dims")
check(dims["final"] == {"rows": len(cleaned), "columns": len(header)}, "report final dims")
check((dims["rows_removed"], dims["rows_quarantined"]) == (len(removed), len(quarantined)), "report removed/quarantined")
for name, rows in (("cleaned", cleaned), ("quarantine", quarantined), ("removed_rows", removed)):
    check(report["outputs"][name]["rows"] == len(rows) and report["outputs"][name]["verified"], f"report outputs.{name}")

# ---- removed rows are exactly the injected copies, each identical to the row it names
copy_source_rows = {i + 1 for i, f in enumerate(flags) if "exact_dup_copy" in f}  # data index i -> source index i-1 -> row i+1
named = []
for row in removed:
    match = re.fullmatch(r"exact duplicate of row (\d+)", row[-1])
    check(match is not None, f"removal reason {row[-1]}")
    if match:
        named.append(int(match.group(1)))
check(set(named) == copy_source_rows, "removed rows are exactly the injected copies")
check(all(raw[n - 2] == raw[n - 1] for n in named), "named rows are identical in the raw input")

# ---- quarantine reasons
reason_ids = Counter(rid for row in quarantined for rid in re.findall(r"\[([a-z_0-9]+)\]", row[-1]))
v = expected["validation"]
wanted = {"id_present", "status_allowed"} if workflow == "customers" else {"order_id_present", "channel_allowed"}
check(reason_ids == Counter({k: v[k] for k in wanted}), f"quarantine reasons {dict(reason_ids)}")

# ---- validation and conversion figures
rules = {r["id"]: r["failures"] for r in report["validation"]["rules"]}
check(rules == expected["validation"], f"validation per rule {rules} vs {expected['validation']}")
check(report["validation"]["issue_count"] == expected["validation_issue_count"], "validation issue count")
check(report["validation"]["rows_with_issues"] == expected["rows_with_issues"], "distinct rows with issues")
by_column = Counter(f["column"] for f in report["conversion_failures"])
date_col, num_col = ("signup_date", "annual_spend") if workflow == "customers" else ("order_date", "quantity")
check(by_column[date_col] == expected["date_failures"], f"date failures {by_column[date_col]}")
check(by_column[num_col] == expected["number_failures"], f"number failures {by_column[num_col]}")
check(report["duplicates"]["keys"]["duplicate_key_groups"] == expected["key_duplicate_groups"], "key duplicate groups")
check(all(len(g["source_rows"]) == 2 for g in report["duplicates"]["keys"]["groups"]), "key groups of 2")

# ---- every value of every kept row, re-derived independently
kept = [i for i, f in enumerate(flags) if "exact_dup_copy" not in f and not (f & quarantine_flags)]
check(len(kept) == len(cleaned), "kept rows")
mismatch = Counter()
if workflow == "customers":
    city_map = {"new york": "New York", "nyc": "New York", "boston": "Boston", "chicago": "Chicago",
                "seattle": "Seattle", "austin": "Austin", "denver": "Denver"}
    for i, out in zip(kept, cleaned):
        r = raw[i]
        city = " ".join(r[3].split())
        spend = r[5].strip()
        if spend == "" or spend.casefold() in {"n/a", "null"}:
            spend = ""
        else:
            number = spend.replace("$", "").strip()
            spend = fmt_float(float(Decimal(number.replace(",", "")))) if re.fullmatch(r"-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?", number) else r[5]
        status = " ".join(r[6].split())
        exp = [
            r[0],
            " ".join(r[1].split()),
            r[2].strip().lower(),
            city_map.get(city.casefold(), city),
            parse_date(r[4], ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]) or r[4],
            spend,
            {"active": "active", "inactive": "inactive", "closed": "closed"}.get(status.casefold(), status),
            r[7],
        ]
        mismatch.update(name for name, a, b in zip(cols, exp, out) if a != b)
    check(all(re.fullmatch(r"\d{6}", row[0]) for row in cleaned), "customer ids keep 6 digits (leading zeros)")
    check(sum(row[0].startswith("0") for row in cleaned) > 0.9 * len(cleaned), "leading-zero ids present in output")
else:
    for i, out, typed in zip(kept, cleaned, cleaned_typed):
        r = raw[i]
        date = excel_text(r[1])
        quantity = excel_text(r[4]).strip()
        channel = excel_text(r[6]).strip()
        exp = [
            excel_text(r[0]),
            parse_date(date, ["%Y-%m-%d", "%d.%m.%Y"]) or date,
            excel_text(r[2]),
            excel_text(r[3]).strip().upper(),
            str(int(quantity)) if quantity.isdigit() else quantity,
            "" if r[5] is None else fmt_float(float(Decimal(str(r[5])))),
            {"online": "Online", "store": "Store", "phone": "Phone"}.get(channel.casefold(), channel),
        ]
        mismatch.update(name for name, a, b in zip(cols, exp, out) if a != b)
        types = {name: t for name, (value, t) in zip(cols, typed) if value is not None}
        if "unit_price" in types and types["unit_price"] != "n":
            mismatch["unit_price_type"] += 1
        if types.get("customer_id") != "s" or types.get("quantity") != "s":
            mismatch["text_column_type"] += 1
    check(sum(1 for row in cleaned if re.fullmatch(r"\d{6}", row[2]) and row[2].startswith("0")) > 0.5 * len(cleaned),
          "leading-zero customer ids present")
check(not mismatch, f"row-by-row mismatches {dict(mismatch)}")

print(f"{workflow} run {run}: {checks} checks, {len(failures)} failed")
for failure in failures:
    print("  FAILED:", failure)
sys.exit(1 if failures else 0)
